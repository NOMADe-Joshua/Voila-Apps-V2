"""
smart_databaser: unit tests for the data_manager model layer.

Covers: ExperimentState process add/remove + renumbering, no-clobber write path,
varying-field scope promotion, and material-gated progress counting.
"""

from __future__ import annotations

from unittest.mock import MagicMock, patch

import pytest
import requests

from perotf_utils.process_specs import (
    build_field_paths,
    build_field_value_multipliers,
)

# App modules come from the conftest's ``mods`` fixture (never imported by bare name); the
# autouse fixture below binds them to these module globals for the tests.
dm = gc = ac = app_module = None


@pytest.fixture(scope="module", autouse=True)
def _bind_app_modules(mods):
    global dm, gc, ac, app_module
    dm = mods["data_manager"]
    gc = mods["gui_components"]
    ac = mods["alias_config"]
    app_module = mods["app"]


# ---------------------------------------------------------------------------
# ExperimentState.add_process / remove_process / renumbering
# ---------------------------------------------------------------------------


def test_add_process_assigns_sequence_index_starting_at_one(fresh_state):
    p1 = fresh_state.add_process("Cleaning UV-Ozone")
    p2 = fresh_state.add_process("Spin Coating")
    assert p1.sequence_index == 1
    assert p2.sequence_index == 2


def test_remove_process_renumbers_remaining(fresh_state):
    fresh_state.add_process("Cleaning UV-Ozone")
    fresh_state.add_process("Spin Coating")
    p3 = fresh_state.add_process("Evaporation")

    fresh_state.remove_process(1)

    assert [p.process_type for p in fresh_state.process_sequence] == [
        "Spin Coating",
        "Evaporation",
    ]
    assert fresh_state.get_process(1).process_type == "Spin Coating"
    assert fresh_state.get_process(2).process_type == "Evaporation"
    assert p3.sequence_index == 2


def test_get_process_missing_raises_keyerror(fresh_state):
    with pytest.raises(KeyError):
        fresh_state.get_process(1)


def test_add_process_at_index_inserts_and_renumbers(fresh_state):
    fresh_state.add_process("Cleaning UV-Ozone")
    fresh_state.add_process("Evaporation")
    fresh_state.add_process("Spin Coating", at_index=1)

    assert [p.process_type for p in fresh_state.process_sequence] == [
        "Cleaning UV-Ozone",
        "Spin Coating",
        "Evaporation",
    ]


# ---------------------------------------------------------------------------
# set_field_if_empty / set_field_manual -- no-clobber path
# ---------------------------------------------------------------------------


def test_set_field_if_empty_writes_when_blank():
    spec = dm.ProcessFieldSpec(key="Material name")
    wrote = dm.set_field_if_empty(
        spec, "PCBM", dm.FieldProvenance(source="batch_template", source_batch_id="B1")
    )
    assert wrote is True
    assert spec.value == "PCBM"
    assert spec.provenance.source == "batch_template"


def test_set_field_if_empty_does_not_overwrite_existing_value():
    spec = dm.ProcessFieldSpec(key="Material name", value="Aluminium")
    wrote = dm.set_field_if_empty(spec, "PCBM")
    assert wrote is False
    assert spec.value == "Aluminium"


def test_set_field_if_empty_treats_empty_string_as_blank():
    spec = dm.ProcessFieldSpec(key="Notes", value="")
    wrote = dm.set_field_if_empty(spec, "filled from history")
    assert wrote is True
    assert spec.value == "filled from history"


def test_set_field_manual_always_overwrites():
    spec = dm.ProcessFieldSpec(key="Material name", value="PCBM", is_outlier=True)
    dm.set_field_manual(spec, "Spiro-OMeTAD")
    assert spec.value == "Spiro-OMeTAD"
    assert spec.provenance.source == "manual"
    assert spec.is_outlier is False


def test_set_field_if_empty_on_varying_field_requires_sample_number():
    spec = dm.ProcessFieldSpec(key="Solvent 1 name", varies=True)
    with pytest.raises(ValueError):
        dm.set_field_if_empty(spec, "DMF")


def test_set_field_if_empty_on_varying_field_writes_per_sample():
    spec = dm.ProcessFieldSpec(key="Solvent 1 name", varies=True)
    wrote_1 = dm.set_field_if_empty(spec, "DMF", sample_number=1)
    wrote_2 = dm.set_field_if_empty(spec, "DMSO", sample_number=1)  # sample 1 now filled
    wrote_3 = dm.set_field_if_empty(spec, "NMP", sample_number=2)

    assert wrote_1 is True
    assert wrote_2 is False  # no-clobber: sample 1 already has "DMF"
    assert wrote_3 is True
    assert spec.per_sample_values == {1: "DMF", 2: "NMP"}


def test_set_field_manual_on_varying_field_overwrites_per_sample():
    spec = dm.ProcessFieldSpec(key="Solvent 1 name", varies=True, per_sample_values={1: "DMF"})
    dm.set_field_manual(spec, "DMSO", sample_number=1)
    assert spec.per_sample_values[1] == "DMSO"
    assert spec.per_sample_provenance[1].source == "manual"


# ---------------------------------------------------------------------------
# set_field_varies -- scope promotion
# ---------------------------------------------------------------------------


def test_set_field_varies_seeds_only_first_sample_from_constant_value():
    spec = dm.ProcessFieldSpec(
        key="Annealing temperature [C]",
        value=120,
        provenance=dm.FieldProvenance(source="batch_template", source_batch_id="B1"),
    )
    dm.set_field_varies(spec, True, sample_numbers=[1, 2, 3])

    assert spec.varies is True
    # Only the first row (matrix row order) is pre-filled - the rest start blank, per
    # the product decision that the table shouldn't look "already filled in" beyond the
    # one row that already had a value before "varies" was checked. The populate-down
    # button (populate_column_from_first) is the explicit way to copy it into the rest.
    assert spec.per_sample_values == {1: 120}
    assert spec.per_sample_provenance[1].source == "batch_template"


def test_set_field_varies_does_not_clobber_existing_first_sample_value():
    spec = dm.ProcessFieldSpec(
        key="Annealing temperature [C]",
        value=120,
        per_sample_values={1: 150},
    )
    dm.set_field_varies(spec, True, sample_numbers=[1, 2, 3])

    assert spec.per_sample_values == {1: 150}


def test_set_field_varies_off_preserves_per_sample_values():
    spec = dm.ProcessFieldSpec(key="X", varies=True, per_sample_values={1: "a"})
    dm.set_field_varies(spec, False, sample_numbers=[1])
    assert spec.varies is False
    assert spec.per_sample_values == {1: "a"}  # not destroyed, can re-enable later


# ---------------------------------------------------------------------------
# Material-gated progress
# ---------------------------------------------------------------------------


def test_non_material_process_always_counts():
    process = dm.ProcessInstance(process_type="Laser Scribing", sequence_index=1)
    assert process.counts_toward_progress() is True


def test_material_gated_process_excluded_when_material_empty():
    process = dm.ProcessInstance(
        process_type="Spin Coating",
        sequence_index=1,
        field_specs={"Material name": dm.ProcessFieldSpec(key="Material name")},
    )
    assert process.counts_toward_progress() is False


def test_material_gated_process_excluded_when_material_field_missing_entirely():
    process = dm.ProcessInstance(process_type="Spin Coating", sequence_index=1)
    assert process.counts_toward_progress() is False


def test_material_gated_process_counts_once_material_filled():
    process = dm.ProcessInstance(
        process_type="Spin Coating",
        sequence_index=1,
        field_specs={"Material name": dm.ProcessFieldSpec(key="Material name", value="PCBM")},
    )
    assert process.counts_toward_progress() is True


def test_material_gated_process_counts_when_material_varies_and_any_sample_filled():
    process = dm.ProcessInstance(
        process_type="Evaporation",
        sequence_index=1,
        field_specs={
            "Material name": dm.ProcessFieldSpec(
                key="Material name", varies=True, per_sample_values={1: "", 2: "Aluminium"}
            )
        },
    )
    assert process.counts_toward_progress() is True


# ---------------------------------------------------------------------------
# process_sequence_to_dicts
# ---------------------------------------------------------------------------


def test_process_sequence_to_dicts_always_leads_with_experiment_info(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 2})
    dicts = dm.process_sequence_to_dicts(fresh_state)
    assert dicts[0] == {"process": "Experiment Info"}
    assert dicts[1] == {"process": "Spin Coating", "config": {"solvents": 2}}


def test_process_sequence_to_dicts_omits_config_key_when_empty(fresh_state):
    fresh_state.add_process("Evaporation")
    dicts = dm.process_sequence_to_dicts(fresh_state)
    assert dicts[1] == {"process": "Evaporation"}


# ---------------------------------------------------------------------------
# generate_header_workbook / build_column_map -- reuses Excel_creator's sheet_experiment.py
# exactly, reconstructs the column map by reading back row 1 / row 2 rather than
# reimplementing generate_steps_for_process (a nested, unexported function).
# ---------------------------------------------------------------------------


def test_build_column_map_experiment_info_only(fresh_state):
    workbook = dm.generate_header_workbook(fresh_state)
    column_map = dm.build_column_map(workbook.active)

    assert column_map[(0, "Date")] == 1
    assert column_map[(0, "Nomad ID")] == 6
    assert column_map[(0, "Variation")] == 7
    assert (0, "Number of pixels") in column_map
    assert (0, "Pixel area") in column_map
    assert (0, "Bottom Cell Name") in column_map


def test_build_column_map_assigns_distinct_sequence_index_per_process(fresh_state):
    fresh_state.add_process("Cleaning UV-Ozone", config={"solvents": 1})
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1, "spinsteps": 1})
    fresh_state.add_process("Evaporation")

    workbook = dm.generate_header_workbook(fresh_state)
    column_map = dm.build_column_map(workbook.active)

    sequence_indices_present = {seq for seq, _ in column_map}
    assert sequence_indices_present == {0, 1, 2, 3}

    material_col_process2 = column_map[(2, "Material name")]
    material_col_process3 = column_map[(3, "Material name")]
    assert material_col_process2 != material_col_process3


def test_build_column_map_repeated_process_type_gets_non_overlapping_columns(fresh_state):
    """A real experiment file has two separate Slot Die Coating processes
    (positions 5 and 6) -- must not collapse into one column range."""
    fresh_state.add_process("Slot Die Coating", config={"solvents": 1, "solutes": 1})
    fresh_state.add_process("Slot Die Coating", config={"solvents": 3, "solutes": 2})

    workbook = dm.generate_header_workbook(fresh_state)
    column_map = dm.build_column_map(workbook.active)

    cols_process1 = {col for (seq, _key), col in column_map.items() if seq == 1}
    cols_process2 = {col for (seq, _key), col in column_map.items() if seq == 2}

    assert cols_process1.isdisjoint(cols_process2)
    # process 2 has more solvents/solutes configured, so more columns
    assert len(cols_process2) > len(cols_process1)
    assert (2, "Solvent 3 name") in column_map
    assert (1, "Solvent 3 name") not in column_map


def test_build_column_map_slot_die_coating_matches_real_file_solvent_solute_shape(fresh_state):
    """A Slot Die Coating with 3 solvents and 2 solutes, checked against the CURRENT
    Excel_creator generator's actual output (peroTF's solvent name/volume and solute
    type/concentration columns, plus its slot-die-specific columns)."""
    fresh_state.add_process("Slot Die Coating", config={"solvents": 3, "solutes": 2})
    workbook = dm.generate_header_workbook(fresh_state)
    column_map = dm.build_column_map(workbook.active)

    for expected_key in [
        "Solvent 1 name",
        "Solvent 2 name",
        "Solvent 3 name",
        "Solvent 3 volume [uL]",
        "Solute 1 type",
        "Solute 2 Concentration [mM]",
        "Coating run",
        "Coated area [mm²]",
    ]:
        assert (1, expected_key) in column_map, expected_key

    assert (1, "Solvent 3 chemical ID") not in column_map
    assert (1, "Room temperature [°C]") not in column_map


def test_sync_field_specs_from_columns_creates_specs_additively(fresh_state):
    fresh_state.add_process("Evaporation")
    column_map = dm.build_column_map(dm.generate_header_workbook(fresh_state).active)

    dm.sync_field_specs_from_columns(fresh_state, column_map)

    assert "Material name" in fresh_state.get_process(1).field_specs
    assert "Date" in fresh_state.experiment_info_fields
    assert "Number of pixels" in fresh_state.experiment_info_fields
    assert isinstance(fresh_state.experiment_info_fields["Number of pixels"], dm.ProcessFieldSpec)


def test_sync_field_specs_from_columns_never_destroys_existing_values(fresh_state):
    process = fresh_state.add_process("Evaporation")
    process.field_specs["Material name"] = dm.ProcessFieldSpec(key="Material name", value="PCBM")

    dm.rebuild_field_specs(fresh_state)

    assert fresh_state.get_process(1).field_specs["Material name"].value == "PCBM"


def test_rebuild_field_specs_returns_column_map_and_syncs(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1, "spinsteps": 1})
    column_map = dm.rebuild_field_specs(fresh_state)

    assert (1, "Rotation speed [rpm]") in column_map
    assert "Rotation speed [rpm]" in fresh_state.get_process(1).field_specs


def test_rebuild_field_specs_multi_spinstep_uses_indexed_rotation_fields(fresh_state):
    """Ties directly to the alias-group scenario planned for step 4: single-step Spin
    Coating uses 'Rotation speed [rpm]', multi-step uses 'Rotation speed {n} [rpm]'."""
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1, "spinsteps": 2})
    column_map = dm.rebuild_field_specs(fresh_state)

    assert (1, "Rotation speed [rpm]") not in column_map
    assert (1, "Rotation speed 1 [rpm]") in column_map
    assert (1, "Rotation speed 2 [rpm]") in column_map


def test_full_real_file_shaped_sequence_produces_disjoint_column_ranges(fresh_state):
    """Coarse structural regression test mirroring the real batch's process shape:
    Experiment Info, Laser Scribing, Cleaning O2-Plasma, Cleaning UV-Ozone, Spin Coating,
    two Slot Die Coating steps, Evaporation, ALD, Laser Scribing, Evaporation, Laser
    Scribing (11 processes total, matching the 11 merged header ranges observed in a
    real experiment file)."""
    fresh_state.add_process("Laser Scribing")
    fresh_state.add_process("Cleaning O2-Plasma", config={"solvents": 2})
    fresh_state.add_process("Cleaning UV-Ozone", config={"solvents": 1})
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1, "spinsteps": 1})
    fresh_state.add_process("Slot Die Coating", config={"solvents": 1, "solutes": 1})
    fresh_state.add_process("Slot Die Coating", config={"solvents": 3, "solutes": 2})
    fresh_state.add_process("Evaporation")
    fresh_state.add_process("ALD")
    fresh_state.add_process("Laser Scribing")
    fresh_state.add_process("Evaporation")
    fresh_state.add_process("Laser Scribing")

    assert len(fresh_state.process_sequence) == 11

    column_map = dm.rebuild_field_specs(fresh_state)
    sequence_indices_present = {seq for seq, _ in column_map}
    assert sequence_indices_present == set(range(12))  # 0..11 (Experiment Info + 11 processes)

    columns_by_sequence: dict[int, set[int]] = {}
    for (seq, _key), col in column_map.items():
        columns_by_sequence.setdefault(seq, set()).add(col)

    all_columns_seen: set[int] = set()
    for seq in range(12):
        cols = columns_by_sequence[seq]
        assert cols.isdisjoint(all_columns_seen)
        all_columns_seen |= cols

    # every process instance got its own field_specs populated
    for process in fresh_state.process_sequence:
        assert process.field_specs, process.process_type


# ---------------------------------------------------------------------------
# ProcessSequenceBuilder -- light smoke tests only (widget construction/interaction): the
# heavy coverage lives in the zero-widget data_manager layer above.
# ---------------------------------------------------------------------------


def test_process_sequence_builder_renders_one_row_per_process(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    fresh_state.add_process("Evaporation")

    builder = gc.ProcessSequenceBuilder(fresh_state)

    assert len(builder.rows_box.children) == 2


def test_process_sequence_builder_add_button_inserts_and_rebuilds_specs(fresh_state):
    fresh_state.add_process("Evaporation")
    builder = gc.ProcessSequenceBuilder(fresh_state)

    builder._add_after(1)

    assert len(fresh_state.process_sequence) == 2
    assert fresh_state.process_sequence[1].process_type == "Generic Process"
    assert len(builder.rows_box.children) == 2


def test_process_sequence_builder_remove_button_deletes_and_renumbers(fresh_state):
    fresh_state.add_process("Evaporation")
    fresh_state.add_process("ALD")
    builder = gc.ProcessSequenceBuilder(fresh_state)

    builder._remove(1)

    assert [p.process_type for p in fresh_state.process_sequence] == ["ALD"]
    assert fresh_state.process_sequence[0].sequence_index == 1
    assert len(builder.rows_box.children) == 1


def test_process_sequence_builder_config_change_updates_state_and_field_specs(fresh_state):
    fresh_state.add_process("Cleaning O2-Plasma", config={"solvents": 1})
    builder = gc.ProcessSequenceBuilder(fresh_state)

    builder._on_config_change(1, "solvents", 3)

    assert fresh_state.get_process(1).config["solvents"] == 3
    assert "Solvent 3" in fresh_state.get_process(1).field_specs


def test_process_sequence_builder_process_type_change_resets_config_and_specs(fresh_state):
    process = fresh_state.add_process("Cleaning O2-Plasma", config={"solvents": 2})
    process.field_specs["Solvent 1"] = dm.ProcessFieldSpec(key="Solvent 1", value="Hellmanex")
    builder = gc.ProcessSequenceBuilder(fresh_state)

    builder._on_process_type_change(1, "Evaporation")

    updated = fresh_state.get_process(1)
    assert updated.process_type == "Evaporation"
    assert "Material name" in updated.field_specs
    assert "Solvent 1" not in updated.field_specs


def test_process_sequence_builder_experiment_info_row_is_fixed_and_first(fresh_state):
    builder = gc.ProcessSequenceBuilder(fresh_state)

    info_row = builder.experiment_info_box.children[0]
    main_row = info_row.children[0]
    # toggle, index label, dropdown, progress label, add button
    _toggle, _index, dropdown, _progress, _add_button = main_row.children

    assert dropdown.value == "Experiment Info"
    assert dropdown.disabled is True


def test_process_sequence_builder_experiment_info_add_button_inserts_first_process(fresh_state):
    builder = gc.ProcessSequenceBuilder(fresh_state)

    builder._add_after(0)

    assert len(fresh_state.process_sequence) == 1
    assert fresh_state.process_sequence[0].process_type == "Generic Process"
    assert fresh_state.process_sequence[0].sequence_index == 1


def test_process_sequence_builder_refresh_picks_up_externally_replaced_sequence(fresh_state):
    """Guards against a real bug: this widget only used to re-render in response to its
    OWN actions, so replacing state.process_sequence from outside (e.g. the
    whole-experiment template picker) left the on-screen rows stale until some unrelated
    edit happened to trigger a re-render."""
    builder = gc.ProcessSequenceBuilder(fresh_state)
    assert len(builder.rows_box.children) == 0

    # Simulates an external mutation (e.g. apply_whole_experiment_template replacing the
    # sequence) - bypasses builder's own add/remove methods, which already trigger a
    # re-render themselves.
    fresh_state.add_process("Evaporation")
    assert len(builder.rows_box.children) == 0  # not yet reflected

    builder.refresh()
    assert len(builder.rows_box.children) == 1


def test_process_sequence_builder_toggle_collapses_and_expands_row(fresh_state):
    fresh_state.add_process("Evaporation")
    builder = gc.ProcessSequenceBuilder(fresh_state)

    assert builder._expanded.get(1, True) is True
    builder._on_toggle(1)
    assert builder._expanded[1] is False
    builder._on_toggle(1)
    assert builder._expanded[1] is True


def test_process_sequence_builder_adopt_section_errors_without_template_batch(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    cache = dm.NomadSessionCache()
    builder = gc.ProcessSequenceBuilder(fresh_state, "url", "token", cache)

    row = builder._process_rows[id(fresh_state.get_process(1))]
    button_row, _caption, _picker_area = row.adopt_section.children
    adopt_button, status = button_row.children
    adopt_button.click()

    assert "template batch" in status.value.lower()


def test_process_sequence_builder_adopt_section_errors_when_batch_missing_process(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    cache = _cache_with("B1", [EVAPORATION_STEP])
    fresh_state.whole_experiment_template_batch_id = "B1"
    builder = gc.ProcessSequenceBuilder(fresh_state, "url", "token", cache)

    row = builder._process_rows[id(fresh_state.get_process(1))]
    button_row, _caption, _picker_area = row.adopt_section.children
    adopt_button, status = button_row.children
    adopt_button.click()

    assert "no" in status.value.lower() and "Spin Coating" in status.value


def test_process_sequence_builder_adopt_section_single_occurrence_applies_directly(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with("B1", [SPIN_COATING_STEP])
    fresh_state.whole_experiment_template_batch_id = "B1"
    builder = gc.ProcessSequenceBuilder(fresh_state, "url", "token", cache)

    row = builder._process_rows[id(fresh_state.get_process(1))]
    button_row, _caption, _picker_area = row.adopt_section.children
    adopt_button, _status = button_row.children
    adopt_button.click()

    process = fresh_state.get_process(1)
    assert process.field_specs["Material name"].value == "Me4PACz"
    assert process.source_override_batch_id == "B1"


def test_process_sequence_builder_adopt_section_multiple_occurrences_shows_material_picker(
    fresh_state,
):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    # Distinct positon_in_experimental_plan required - see distinct_steps_for_process_type.
    second_step = {
        **SPIN_COATING_STEP,
        "positon_in_experimental_plan": 9.0,
        "layer": [{"layer_material_name": "SecondLayerMaterial"}],
    }
    cache = _cache_with("B1", [SPIN_COATING_STEP, second_step])
    fresh_state.whole_experiment_template_batch_id = "B1"
    builder = gc.ProcessSequenceBuilder(fresh_state, "url", "token", cache)

    row = builder._process_rows[id(fresh_state.get_process(1))]
    button_row, _caption, picker_area = row.adopt_section.children
    adopt_button, _status = button_row.children
    adopt_button.click()

    occurrence_dropdown, confirm_button = picker_area.children
    occurrence_dropdown.value = 1
    confirm_button.click()

    assert fresh_state.get_process(1).field_specs["Material name"].value == "SecondLayerMaterial"


# ---------------------------------------------------------------------------
# NOMAD live value sourcing -- fixtures below are shaped exactly like real archive
# data pulled live from a real batch (steps "spin coating Me4PACz" /
# "evaporation C60") during implementation, trimmed to the fields the mapping uses.
# Tests here mock the network layer -- they never hit the real server.
# ---------------------------------------------------------------------------

SPIN_COATING_STEP = {
    "method": "Spin Coating",
    "name": "spin coating Me4PACz",
    "positon_in_experimental_plan": 3.0,
    "samples": [{"lab_id": "KIT_JoDa_20260526_12_1_1"}, {"lab_id": "KIT_JoDa_20260526_12_1_2"}],
    "layer": [{"layer_type": "Hole Transport Layer", "layer_material_name": "Me4PACz"}],
    "solution": [
        {
            "solution_volume": 0.12,
            "solution_details": {
                "solute": [{"name": "Me4PACz", "concentration_mol": 0.003}],
                "solvent": [{"name": "Ethanol 0.12 milliliter", "chemical_2": {"name": "Ethanol"}}],
            },
        }
    ],
    "annealing": {"temperature": 100.0, "time": 600.0},
    "recipe_steps": [{"time": 30.0, "speed": 3000.0}],
}

SPIN_COATING_STEP_VARIATION = {
    **SPIN_COATING_STEP,
    "samples": [{"lab_id": "KIT_JoDa_20260526_12_2_1"}],
    "annealing": {"temperature": 130.0, "time": 600.0},
}

SPIN_COATING_STEP_CLOSE = {
    **SPIN_COATING_STEP,
    "samples": [{"lab_id": "KIT_JoDa_20260526_12_3_1"}],
    "annealing": {"temperature": 102.0, "time": 600.0},
}

SPIN_COATING_STEP_EXTREME_OUTLIER = {
    **SPIN_COATING_STEP,
    "samples": [{"lab_id": "KIT_JoDa_20260526_12_4_1"}],
    "annealing": {"temperature": 99999.0, "time": 600.0},
}

EVAPORATION_STEP = {
    "method": "Evaporation",
    "name": "evaporation C60",
    "positon_in_experimental_plan": 7.0,
    "samples": [{"lab_id": "KIT_JoDa_20260526_12_1_1"}],
    "layer": [{"layer_type": "Electron Transport Layer", "layer_material_name": "C60"}],
    "organic_evaporation": [{"thickness": 40.0, "start_rate": 0.5}],
}


def test_steps_for_process_type_filters_by_method():
    steps = [SPIN_COATING_STEP, EVAPORATION_STEP]
    assert dm.steps_for_process_type(steps, "Evaporation") == [EVAPORATION_STEP]
    assert dm.steps_for_process_type(steps, "ALD") == []


def test_fetch_process_field_values_returns_mapped_values_and_source_sample():
    cache = dm.NomadSessionCache()
    cache._processing_steps_by_batch["B1"] = [SPIN_COATING_STEP]

    values, source_sample_id = dm.fetch_process_field_values(
        "url", "token", cache, "B1", "Spin Coating"
    )

    assert values["Material name"] == "Me4PACz"
    assert values["Layer type"] == "Hole Transport Layer"
    assert values["Solvent 1 name"] == "Ethanol"
    assert values["Solute 1 type"] == "Me4PACz"
    assert values["Solute 1 Concentration [mM]"] == 0.003
    assert values["Rotation speed [rpm]"] == 3000.0
    assert values["Annealing temperature [°C]"] == 100.0
    assert source_sample_id == "KIT_JoDa_20260526_12_1_1"


def test_fetch_process_field_values_unmapped_process_type_returns_empty():
    cache = dm.NomadSessionCache()
    cache._processing_steps_by_batch["B1"] = [SPIN_COATING_STEP]
    values, source = dm.fetch_process_field_values("url", "token", cache, "B1", "Laser Scribing")
    assert values == {}
    assert source is None


def test_fetch_process_field_values_missing_occurrence_returns_empty():
    cache = dm.NomadSessionCache()
    cache._processing_steps_by_batch["B1"] = [SPIN_COATING_STEP]
    values, source = dm.fetch_process_field_values(
        "url", "token", cache, "B1", "Spin Coating", occurrence=1
    )
    assert values == {}
    assert source is None


def test_nomad_session_cache_calls_api_once_per_batch():
    cache = dm.NomadSessionCache()
    with (
        patch.object(dm, "get_ids_in_batch", return_value=["S1", "S2"]) as mock_ids,
        patch.object(dm, "get_processing_steps", return_value=[SPIN_COATING_STEP]) as mock_steps,
    ):
        first = cache.get_processing_steps("url", "token", "B1")
        second = cache.get_processing_steps("url", "token", "B1")

        assert first == second == [SPIN_COATING_STEP]
        mock_ids.assert_called_once_with("url", "token", ["B1"])
        mock_steps.assert_called_once_with("url", "token", ["S1", "S2"])


def test_nomad_session_cache_clear_forces_refetch():
    cache = dm.NomadSessionCache()
    with (
        patch.object(dm, "get_ids_in_batch", return_value=["S1"]),
        patch.object(dm, "get_processing_steps", return_value=[]) as mock_steps,
    ):
        cache.get_processing_steps("url", "token", "B1")
        cache.clear()
        cache.get_processing_steps("url", "token", "B1")

        assert mock_steps.call_count == 2


def test_autofill_process_from_batch_writes_only_existing_specs_no_clobber(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    # pre-fill one field manually -- must survive autofill untouched
    process.field_specs["Material name"].value = "manually chosen material"

    cache = dm.NomadSessionCache()
    cache._processing_steps_by_batch["B1"] = [SPIN_COATING_STEP]

    written = dm.autofill_process_from_batch(process, "url", "token", cache, "B1")

    assert process.field_specs["Material name"].value == "manually chosen material"
    assert process.field_specs["Layer type"].value == "Hole Transport Layer"
    assert process.field_specs["Layer type"].provenance.source == "batch_template"
    assert process.field_specs["Layer type"].provenance.source_batch_id == "B1"
    assert (
        process.field_specs["Layer type"].provenance.source_sample_id == "KIT_JoDa_20260526_12_1_1"
    )
    assert written > 0


def test_autofill_process_from_batch_skips_fields_with_no_spec(fresh_state):
    """Evaporation's field_specs come from generate_header_workbook, which won't include
    every possible archive field -- autofill must not crash on unmapped-but-fetched keys."""
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)

    cache = dm.NomadSessionCache()
    cache._processing_steps_by_batch["B1"] = [EVAPORATION_STEP]

    written = dm.autofill_process_from_batch(process, "url", "token", cache, "B1")

    assert process.field_specs["Material name"].value == "C60"
    assert written > 0


def test_compute_field_distribution_for_occurrence_scopes_to_same_position():
    cache = dm.NomadSessionCache()
    cache._processing_steps_by_batch["B1"] = [SPIN_COATING_STEP, SPIN_COATING_STEP_VARIATION]

    distribution = dm.compute_field_distribution_for_occurrence(
        "url", "token", cache, "B1", "Spin Coating", "Annealing temperature [°C]"
    )

    assert distribution == [100.0, 130.0]


def test_compute_field_distribution_for_occurrence_excludes_different_position():
    other_position_step = {**SPIN_COATING_STEP_VARIATION, "positon_in_experimental_plan": 9.0}
    cache = dm.NomadSessionCache()
    cache._processing_steps_by_batch["B1"] = [SPIN_COATING_STEP, other_position_step]

    distribution = dm.compute_field_distribution_for_occurrence(
        "url", "token", cache, "B1", "Spin Coating", "Annealing temperature [°C]"
    )

    assert distribution == [100.0]


def test_compute_field_distribution_for_occurrence_exclude_occurrence_drops_own_step():
    """SPIN_COATING_STEP/_CLOSE/_EXTREME_OUTLIER all share one positon_in_experimental_plan
    (representing different samples at the same sequence position), so this is a single
    distinct occurrence (0) - occurrence indexes into DISTINCT positions, see
    distinct_steps_for_process_type; the "own step" that exclude_occurrence drops is
    whichever raw step distinct_steps_for_process_type picked as that position's
    representative (the first-listed one, SPIN_COATING_STEP itself, temp 100.0)."""
    cache = dm.NomadSessionCache()
    cache._processing_steps_by_batch["B1"] = [
        SPIN_COATING_STEP,
        SPIN_COATING_STEP_CLOSE,
        SPIN_COATING_STEP_EXTREME_OUTLIER,
    ]

    with_self = dm.compute_field_distribution_for_occurrence(
        "url", "token", cache, "B1", "Spin Coating", "Annealing temperature [°C]", occurrence=0
    )
    without_self = dm.compute_field_distribution_for_occurrence(
        "url",
        "token",
        cache,
        "B1",
        "Spin Coating",
        "Annealing temperature [°C]",
        occurrence=0,
        exclude_occurrence=True,
    )

    assert with_self == [100.0, 102.0, 99999.0]
    assert without_self == [102.0, 99999.0]


def test_self_inclusive_distribution_masks_extreme_outlier_regression():
    """Documents why autofill_process_from_batch uses exclude_occurrence=True: with the
    candidate value included in its own reference population, a single extreme point
    inflates the population stdev enough that even a wildly extreme value never exceeds a
    fixed z-score threshold at small n - the outlier masks itself."""
    self_inclusive_distribution = [100.0, 102.0, 99999.0]
    assert dm.is_outlier(99999.0, self_inclusive_distribution) is False  # masked

    leave_one_out_distribution = [100.0, 102.0]
    assert dm.is_outlier(99999.0, leave_one_out_distribution) is True  # correctly flagged


def test_is_outlier_flags_value_far_from_distribution():
    distribution = [100.0, 102.0, 98.0, 101.0]
    assert dm.is_outlier(200.0, distribution) is True
    assert dm.is_outlier(100.5, distribution) is False


def test_is_outlier_requires_at_least_two_points():
    assert dm.is_outlier(100.0, [50.0]) is False
    assert dm.is_outlier(100.0, []) is False


def test_is_outlier_ignores_non_numeric_value():
    assert dm.is_outlier("not a number", [1.0, 2.0, 3.0]) is False


# ---------------------------------------------------------------------------
# Field-mapping source (shared/perotf_utils/process_specs.py) -- unit_verified flags
# and new process types/fields are meant to be editable there without touching
# data_manager.py. build_field_paths() is the module-level function under test; passing
# it a small synthetic PROCESSES-shaped dict (instead of the real one) isolates each
# mechanism the same way a temp field_mappings.json used to before this module existed.
# ---------------------------------------------------------------------------


def test_process_type_field_paths_is_the_loaded_source():
    assert build_field_paths() == dm.PROCESS_TYPE_FIELD_PATHS


def test_process_type_field_paths_matches_spec_unit_verified_flags():
    from perotf_utils.process_specs import PROCESSES

    spin_coating_fields = PROCESSES["Spin Coating"]["fields"]
    for excel_key, field_spec in spin_coating_fields.items():
        if "path" not in field_spec and "paths" not in field_spec:
            continue
        _path, unit_verified = dm.PROCESS_TYPE_FIELD_PATHS["Spin Coating"][excel_key]
        assert unit_verified == field_spec.get("unit_verified", True), excel_key


def test_build_field_paths_resolves_indexed_field_templates():
    processes = {
        "Test Process": {
            "fields": {
                "Material name": {"test": "x", "path": ["layer", 0, "layer_material_name"]},
            },
            "indexed": {
                "solvents": [
                    {
                        "excel_key": "Solvent {n} name",
                        "test": lambda n: "x",
                        "path": ["solution", 0, "solvent", "{i}", "name"],
                        "mapping_range": (1, 3),
                    }
                ]
            },
        }
    }

    mappings = build_field_paths(processes)

    # Every entry is normalized to a list of alternative paths (see _get_path_any),
    # even single-path ones like these.
    assert mappings["Test Process"]["Material name"] == (
        [["layer", 0, "layer_material_name"]],
        True,
    )
    assert mappings["Test Process"]["Solvent 1 name"] == (
        [["solution", 0, "solvent", 0, "name"]],
        True,
    )
    assert mappings["Test Process"]["Solvent 3 name"] == (
        [["solution", 0, "solvent", 2, "name"]],
        True,
    )
    assert "Solvent 4 name" not in mappings["Test Process"]


def test_build_field_paths_paths_plural_normalizes_to_alternative_list():
    """'paths' (plural) is for fields whose archive value lives under a different
    parent key depending on other data on the same step (e.g. Evaporation's
    organic_evaporation vs inorganic_evaporation split) - same field name, different
    list. Resolution order is verified end-to-end via
    test_fetch_process_field_values_evaporation_falls_back_across_organic_inorganic
    below, against the real spec."""
    processes = {
        "Evaporation": {
            "fields": {
                "Thickness [nm]": {
                    "test": 1,
                    "paths": [
                        ["organic_evaporation", 0, "thickness"],
                        ["inorganic_evaporation", 0, "thickness"],
                    ],
                    "unit_verified": False,
                }
            },
        }
    }

    mappings = build_field_paths(processes)

    assert mappings["Evaporation"]["Thickness [nm]"] == (
        [["organic_evaporation", 0, "thickness"], ["inorganic_evaporation", 0, "thickness"]],
        False,
    )


def test_build_field_paths_adding_a_field_requires_no_code_change():
    """Proves the 'easy to add/remove' property: a brand-new process type with a field
    marked unit_verified=False round-trips correctly through fetch_process_field_values
    without any change to data_manager.py."""
    processes = {
        "Annealing": {
            "fields": {
                "Annealing temperature [°C]": {
                    "test": 1,
                    "path": ["annealing", "temperature"],
                    "unit_verified": False,
                }
            },
        }
    }

    mappings = build_field_paths(processes)

    assert mappings["Annealing"]["Annealing temperature [°C]"] == (
        [["annealing", "temperature"]],
        False,
    )


def test_build_field_paths_removing_a_field_from_spec_removes_it_from_mapping():
    processes = {
        "Spin Coating": {
            "fields": {
                "Material name": {"test": "x", "path": ["layer", 0, "layer_material_name"]},
            },
        }
    }

    mappings = build_field_paths(processes)

    assert "Material name" in mappings["Spin Coating"]
    assert "Layer type" not in mappings["Spin Coating"]


# ---------------------------------------------------------------------------
# occurrence_index_for_process -- Nth same-type process in the sequence sources from the
# batch's Nth same-type step, not always the first.
# ---------------------------------------------------------------------------


def test_occurrence_index_for_process_first_of_type_is_zero(fresh_state):
    process = fresh_state.add_process("Slot Die Coating")
    assert dm.occurrence_index_for_process(fresh_state, process) == 0


def test_occurrence_index_for_process_second_of_type_is_one(fresh_state):
    fresh_state.add_process("Slot Die Coating")
    fresh_state.add_process("Evaporation")
    second_slot_die = fresh_state.add_process("Slot Die Coating")
    assert dm.occurrence_index_for_process(fresh_state, second_slot_die) == 1


# ---------------------------------------------------------------------------
# clear_autofilled_value -- never touches manual edits
# ---------------------------------------------------------------------------


def test_clear_autofilled_value_clears_matching_source():
    spec = dm.ProcessFieldSpec(
        key="Material name",
        value="PCBM",
        provenance=dm.FieldProvenance(source="batch_template", source_batch_id="B1"),
    )
    dm.clear_autofilled_value(spec, {"batch_template"})
    assert spec.value is None
    assert spec.provenance is None


def test_clear_autofilled_value_never_clears_manual_edit():
    spec = dm.ProcessFieldSpec(
        key="Material name", value="hand-picked", provenance=dm.FieldProvenance(source="manual")
    )
    dm.clear_autofilled_value(spec, {"batch_template", "process_override"})
    assert spec.value == "hand-picked"


def test_clear_autofilled_value_on_varying_field_only_clears_matching_samples():
    spec = dm.ProcessFieldSpec(
        key="Solvent 1 name",
        varies=True,
        per_sample_values={1: "DMF", 2: "manually chosen"},
        per_sample_provenance={
            1: dm.FieldProvenance(source="batch_template", source_batch_id="B1"),
            2: dm.FieldProvenance(source="manual"),
        },
    )
    dm.clear_autofilled_value(spec, {"batch_template"})
    assert spec.per_sample_values == {2: "manually chosen"}


# ---------------------------------------------------------------------------
# apply_whole_experiment_template / apply_process_override / clear_process_override --
# override composition ("last selection at a given scope wins, and scopes nest")
# ---------------------------------------------------------------------------


def _cache_with(
    batch_id: str, steps: list[dict], experiment_info_source: dict | None = None
) -> dm.NomadSessionCache:
    """experiment_info_source defaults to None (pre-cached, not fetched) so tests that
    don't care about Experiment Info autofill never trigger a real network call via
    apply_whole_experiment_template's autofill_experiment_info_from_batch step - pass it
    explicitly to exercise that behavior."""
    cache = dm.NomadSessionCache()
    cache._processing_steps_by_batch[batch_id] = steps
    cache._experiment_info_source_by_batch[batch_id] = experiment_info_source
    return cache


def test_apply_whole_experiment_template_replaces_sequence_and_fills_values(fresh_state):
    """ "Replicate Experiment" replicates the batch's own N steps into an N-process
    sequence and fills them, per the product decision - it does not merely fill values
    into whatever the user had already built."""
    cache = _cache_with("B1", [SPIN_COATING_STEP, EVAPORATION_STEP])

    written = dm.apply_whole_experiment_template(fresh_state, "url", "token", cache, "B1")

    assert fresh_state.whole_experiment_template_batch_id == "B1"
    assert [p.process_type for p in fresh_state.process_sequence] == ["Spin Coating", "Evaporation"]
    assert fresh_state.get_process(1).field_specs["Material name"].value == "Me4PACz"
    assert fresh_state.get_process(2).field_specs["Material name"].value == "C60"
    assert written[1] > 0 and written[2] > 0


def test_apply_whole_experiment_template_discards_prior_sequence_and_manual_edits(fresh_state):
    """Re-picking a (possibly different) template batch REPLACES the whole sequence -
    any processes/overrides/manual edits from before the pick are discarded, same as any
    other re-pick. This is a deliberate product decision, not an oversight."""
    process = fresh_state.add_process("Cleaning UV-Ozone")
    process.field_specs["Notes"] = dm.ProcessFieldSpec(key="Notes", value="manually written note")
    old_cache = _cache_with("OLD_BATCH", [SPIN_COATING_STEP])
    dm.apply_whole_experiment_template(fresh_state, "url", "token", old_cache, "OLD_BATCH")

    new_step = {**SPIN_COATING_STEP, "layer": [{"layer_material_name": "DifferentMaterial"}]}
    new_cache = _cache_with("NEW_BATCH", [new_step])
    dm.apply_whole_experiment_template(fresh_state, "url", "token", new_cache, "NEW_BATCH")

    assert len(fresh_state.process_sequence) == 1
    assert fresh_state.get_process(1).field_specs["Material name"].value == "DifferentMaterial"
    assert fresh_state.whole_experiment_template_batch_id == "NEW_BATCH"
    # the manually-edited "Cleaning UV-Ozone" process from before the first pick is gone
    assert "Cleaning UV-Ozone" not in [p.process_type for p in fresh_state.process_sequence]


def test_apply_process_override_only_touches_target_process(fresh_state):
    template_cache = _cache_with("WHOLE", [SPIN_COATING_STEP, EVAPORATION_STEP])
    dm.apply_whole_experiment_template(fresh_state, "url", "token", template_cache, "WHOLE")

    override_step = {**SPIN_COATING_STEP, "layer": [{"layer_material_name": "OverrideMaterial"}]}
    override_cache = _cache_with("OVERRIDE", [override_step])
    process = fresh_state.get_process(1)
    dm.apply_process_override(fresh_state, process, "url", "token", override_cache, "OVERRIDE")

    assert process.field_specs["Material name"].value == "OverrideMaterial"
    assert process.source_override_batch_id == "OVERRIDE"
    assert process.field_specs["Material name"].provenance.source == "process_override"
    # untouched sibling process stays tied to the whole-experiment template
    assert fresh_state.get_process(2).field_specs["Material name"].value == "C60"


def test_clear_process_override_clears_override_sourced_values(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    override_cache = _cache_with("OVERRIDE", [SPIN_COATING_STEP])
    dm.apply_process_override(fresh_state, process, "url", "token", override_cache, "OVERRIDE")

    dm.clear_process_override(fresh_state, process)

    assert process.source_override_batch_id is None
    assert process.field_specs["Material name"].value is None


# ---------------------------------------------------------------------------
# build_process_sequence_from_batch / infer_config_from_source_step /
# expand_process_config_for_source -- "Replicate Experiment" replicating a whole
# sequence, and both that path and "adopt from template batch" capturing every value a
# source step has (not just as many as the target's current config count allows).
# ---------------------------------------------------------------------------

SPIN_COATING_STEP_TWO_SOLVENTS = {
    "method": "Spin Coating",
    "positon_in_experimental_plan": 1.0,
    "samples": [{"lab_id": "S1"}],
    "layer": [{"layer_type": "HTL", "layer_material_name": "Me4PACz"}],
    "solution": [
        {
            "solution_volume": 0.1,
            "solution_details": {
                "solute": [
                    {"name": "A", "concentration_mol": 1},
                    {"name": "B", "concentration_mol": 2},
                ],
                "solvent": [
                    {"name": "Ethanol 0.1 milliliter", "chemical_2": {"name": "Ethanol"}},
                    {"name": "Water 0.1 milliliter", "chemical_2": {"name": "Water"}},
                ],
            },
        }
    ],
    "recipe_steps": [{"time": 30.0, "speed": 3000.0}],
}


def test_build_process_sequence_from_batch_dedupes_same_position_variations(fresh_state):
    """SPIN_COATING_STEP and SPIN_COATING_STEP_VARIATION share positon_in_experimental_plan
    3.0 (different variation-group annealing temps at the same sequence step) - they
    collapse to ONE ProcessInstance, not two."""
    cache = _cache_with("B1", [SPIN_COATING_STEP, SPIN_COATING_STEP_VARIATION])
    sequence = dm.build_process_sequence_from_batch("url", "token", cache, "B1")
    assert len(sequence) == 1
    assert sequence[0].process_type == "Spin Coating"


def test_build_process_sequence_from_batch_orders_by_position(fresh_state):
    cache = _cache_with("B1", [SPIN_COATING_STEP, EVAPORATION_STEP])  # positions 3.0, 7.0
    sequence = dm.build_process_sequence_from_batch("url", "token", cache, "B1")
    assert [p.process_type for p in sequence] == ["Spin Coating", "Evaporation"]
    assert [p.sequence_index for p in sequence] == [1, 2]


def test_build_process_sequence_from_batch_skips_unknown_process_types(fresh_state):
    unknown_step = {"method": "Some Future Process", "positon_in_experimental_plan": 1.0}
    cache = _cache_with("B1", [unknown_step, EVAPORATION_STEP])
    sequence = dm.build_process_sequence_from_batch("url", "token", cache, "B1")
    assert [p.process_type for p in sequence] == ["Evaporation"]


def test_build_process_sequence_from_batch_sizes_config_from_source_data(fresh_state):
    cache = _cache_with("B1", [SPIN_COATING_STEP_TWO_SOLVENTS])
    sequence = dm.build_process_sequence_from_batch("url", "token", cache, "B1")
    assert sequence[0].config["solvents"] == 2
    assert sequence[0].config["solutes"] == 2


def test_infer_config_from_source_step_detects_solvent_and_solute_counts():
    inferred = dm.infer_config_from_source_step("Spin Coating", SPIN_COATING_STEP_TWO_SOLVENTS)
    assert inferred == {"solvents": 2, "solutes": 2, "spinsteps": 1}


def test_infer_config_from_source_step_unmapped_process_type_returns_empty():
    assert dm.infer_config_from_source_step("Laser Scribing", SPIN_COATING_STEP_TWO_SOLVENTS) == {}


def test_expand_process_config_for_source_widens_config_and_adds_field_specs(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    assert "Solvent 2 name" not in process.field_specs
    cache = _cache_with("B1", [SPIN_COATING_STEP_TWO_SOLVENTS])

    dm.expand_process_config_for_source(fresh_state, process, "url", "token", cache, "B1", 0)

    assert process.config["solvents"] == 2
    assert "Solvent 2 name" in process.field_specs


def test_expand_process_config_for_source_never_shrinks_existing_config(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 5, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with("B1", [SPIN_COATING_STEP_TWO_SOLVENTS])  # only 2 solvents in the source

    dm.expand_process_config_for_source(fresh_state, process, "url", "token", cache, "B1", 0)

    assert process.config["solvents"] == 5


def test_apply_process_override_captures_values_beyond_original_config(fresh_state):
    """Regression test: previously, adopting a source step with more solvents/solutes
    than the target process was configured for silently dropped the extra values."""
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with("B1", [SPIN_COATING_STEP_TWO_SOLVENTS])

    dm.apply_process_override(fresh_state, process, "url", "token", cache, "B1")

    assert process.field_specs["Solvent 1 name"].value == "Ethanol"
    assert process.field_specs["Solvent 2 name"].value == "Water"
    assert process.field_specs["Solute 2 type"].value == "B"


def test_apply_process_override_explicit_occurrence_overrides_positional_default(fresh_state):
    """The 'adopt from template batch' picker lets the user pick a specific occurrence
    (e.g. by material) rather than always taking the positionally-matched one."""
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    # A distinct positon_in_experimental_plan is required - occurrence indexes into
    # DISTINCT sequence positions (see distinct_steps_for_process_type), not raw steps.
    second_step = {
        **SPIN_COATING_STEP,
        "positon_in_experimental_plan": 9.0,
        "layer": [{"layer_material_name": "SecondLayerMaterial"}],
    }
    cache = _cache_with("B1", [SPIN_COATING_STEP, second_step])

    dm.apply_process_override(fresh_state, process, "url", "token", cache, "B1", occurrence=1)

    assert process.field_specs["Material name"].value == "SecondLayerMaterial"


# ---------------------------------------------------------------------------
# list_process_occurrences -- per-occurrence labels for the "adopt from template batch"
# picker when a source batch has multiple same-type steps.
# ---------------------------------------------------------------------------


def test_list_process_occurrences_labels_by_material_when_mapped():
    # SPIN_COATING_STEP_VARIATION shares SPIN_COATING_STEP's positon_in_experimental_plan
    # (representing another sample at the SAME sequence position) so is not a second
    # occurrence here - occurrences are distinct sequence positions, see
    # distinct_steps_for_process_type. Use a genuinely later position instead.
    other_position_step = {**SPIN_COATING_STEP_VARIATION, "positon_in_experimental_plan": 9.0}
    cache = _cache_with("B1", [SPIN_COATING_STEP, other_position_step])
    occurrences = dm.list_process_occurrences("url", "token", cache, "B1", "Spin Coating")
    assert occurrences == [(0, "Me4PACz"), (1, "Me4PACz")]


def test_list_process_occurrences_falls_back_to_generic_label_when_unmapped():
    cache = _cache_with("B1", [{"method": "Laser Scribing", "positon_in_experimental_plan": 1.0}])
    occurrences = dm.list_process_occurrences("url", "token", cache, "B1", "Laser Scribing")
    assert occurrences == [(0, "Occurrence 1")]


def test_list_process_occurrences_empty_when_batch_has_no_such_process():
    cache = _cache_with("B1", [EVAPORATION_STEP])
    assert dm.list_process_occurrences("url", "token", cache, "B1", "Spin Coating") == []


# ---------------------------------------------------------------------------
# preview_value_for_field -- cache-only "what would this field become" hint for the
# greyed-out placeholder in the value input, never triggers a network call.
# ---------------------------------------------------------------------------


def test_preview_value_for_field_reads_from_whole_experiment_template(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    fresh_state.whole_experiment_template_batch_id = "B1"
    cache = _cache_with("B1", [SPIN_COATING_STEP])

    assert dm.preview_value_for_field(fresh_state, process, "Material name", cache) == "Me4PACz"


def test_preview_value_for_field_none_without_active_source(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    preview = dm.preview_value_for_field(
        fresh_state, process, "Material name", dm.NomadSessionCache()
    )
    assert preview is None


def test_preview_value_for_field_none_when_steps_not_cached_yet(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    fresh_state.whole_experiment_template_batch_id = "B1"  # never fetched into this cache
    preview = dm.preview_value_for_field(
        fresh_state, process, "Material name", dm.NomadSessionCache()
    )
    assert preview is None


# ---------------------------------------------------------------------------
# Varying-fields matrix / Variation column -- experiment-wide, no-clobber, never
# autofilled from a template.
# ---------------------------------------------------------------------------


def test_iter_varying_fields_excludes_variation_and_non_varying_fields(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    fresh_state.experiment_info_fields["Variation"].varies = True  # should never appear
    fresh_state.get_process(1).field_specs["Material name"].varies = True

    labels = [label for label, _spec in dm.iter_varying_fields(fresh_state)]

    assert any("Material name" in label for label in labels)
    assert not any("Variation" in label for label in labels)


def test_compute_variation_label_joins_checked_fields_with_delimiter(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.get_process(1)
    dm.set_field_varies(process.field_specs["Material name"], True, [])
    dm.set_field_varies(process.field_specs["Layer type"], True, [])
    process.field_specs["Material name"].per_sample_values[1] = "PCBM"
    process.field_specs["Layer type"].per_sample_values[1] = "ETL"

    label = dm.compute_variation_label(fresh_state, sample_number=1)

    assert label == "material-name-PCBM_layer-type-ETL"


def test_compute_variation_label_skips_unfilled_varying_fields(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.get_process(1)
    dm.set_field_varies(process.field_specs["Material name"], True, [])
    process.field_specs["Material name"].per_sample_values[1] = "PCBM"
    dm.set_field_varies(process.field_specs["Layer type"], True, [])
    # Layer type left empty for sample 1

    label = dm.compute_variation_label(fresh_state, sample_number=1)

    assert label == "material-name-PCBM"


def test_update_variation_column_writes_empty_slots_only(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    fresh_state.add_sample(variation_group_index=0, sample_number=2)
    process = fresh_state.get_process(1)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())
    process.field_specs["Material name"].per_sample_values = {1: "PCBM", 2: "Spiro"}
    # sample 1's Variation already has a value (manual or prior computation) - must survive
    fresh_state.experiment_info_fields["Variation"].varies = True
    fresh_state.experiment_info_fields["Variation"].per_sample_values[1] = "user set this"

    written = dm.update_variation_column(fresh_state)

    variation_spec = fresh_state.experiment_info_fields["Variation"]
    assert variation_spec.per_sample_values[1] == "user set this"
    assert variation_spec.per_sample_values[2] == "material-name-Spiro"
    assert written == 1


def test_update_variation_column_never_sourced_from_batch_template(fresh_state):
    """Variation must never be autofilled/inherited from a template - confirmed by the
    fact autofill_process_from_batch only ever touches ProcessInstance.field_specs, never
    experiment_info_fields, so Variation can't appear in a template's written fields."""
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with("B1", [SPIN_COATING_STEP])

    dm.apply_whole_experiment_template(fresh_state, "url", "token", cache, "B1")

    assert "Variation" not in fresh_state.get_process(1).field_specs
    assert fresh_state.experiment_info_fields["Variation"].value is None
    assert fresh_state.experiment_info_fields["Variation"].per_sample_values == {}


# ---------------------------------------------------------------------------
# alias_config.resolve_progress_units -- scoped per process instance, config-driven
# (config/alias_groups.json), primary scenario: single- vs multi-step Spin Coating
# rotation-speed field naming.
# ---------------------------------------------------------------------------


def test_resolve_progress_units_groups_single_vs_multistep_rotation_speed():
    field_keys = ["Material name", "Rotation speed [rpm]"]
    units = ac.resolve_progress_units("Spin Coating", field_keys)
    assert ["Material name"] in units
    assert ["Rotation speed [rpm]"] in units


def test_resolve_progress_units_merges_indexed_variants_into_one_unit():
    field_keys = ["Rotation speed 1 [rpm]", "Rotation speed 2 [rpm]"]
    units = ac.resolve_progress_units("Spin Coating", field_keys)
    assert len(units) == 1
    assert set(units[0]) == {"Rotation speed 1 [rpm]", "Rotation speed 2 [rpm]"}


def test_resolve_progress_units_does_not_cross_process_type_boundary():
    """A field pattern only matches within its declared process_type - Evaporation has no
    alias group, so its fields are never merged even if named similarly."""
    field_keys = ["Rotation speed [rpm]"]
    units = ac.resolve_progress_units("Evaporation", field_keys)
    assert units == [["Rotation speed [rpm]"]]


def test_resolve_progress_units_with_custom_alias_groups_config():
    custom_groups = [
        {
            "id": "thickness_alias",
            "members": [
                {"process_type": "ALD", "field_pattern": "Thickness [nm]"},
                {"process_type": "ALD", "field_pattern": "Film thickness [nm]"},
            ],
        }
    ]
    units = ac.resolve_progress_units(
        "ALD", ["Thickness [nm]", "Film thickness [nm]"], alias_groups=custom_groups
    )
    assert len(units) == 1


# ---------------------------------------------------------------------------
# Material-gated progress bar
# ---------------------------------------------------------------------------


def test_compute_process_progress_excludes_ungated_material_process(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    filled, total = dm.compute_process_progress(process)
    assert (filled, total) == (0, 0)


def test_compute_process_progress_counts_after_material_filled(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    process.field_specs["Material name"].value = "PCBM"

    filled, total = dm.compute_process_progress(process)

    assert filled == 1  # only "Material name" itself is filled
    assert total > 1  # other fields exist and count toward the denominator now


def test_compute_process_progress_non_material_process_always_counted(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    filled, total = dm.compute_process_progress(process)
    assert total > 0
    assert filled == 0


def test_compute_process_progress_merges_alias_group_as_one_unit(fresh_state):
    """A single real Spin Coating instance only ever generates ONE rotation-speed naming
    variant (spinsteps picks exactly one branch in sheet_experiment.py), so the two
    variants never naturally coexist in one instance's field_specs. This test injects the
    second variant synthetically to exercise the alias-merge mechanism directly, since the
    config schema is meant to be ready for whatever real analogous-field case comes up
    (typos, schema drift, etc.), not only the motivating spinsteps example."""
    process = fresh_state.add_process(
        "Spin Coating", config={"solvents": 0, "solutes": 0, "spinsteps": 1}
    )
    dm.rebuild_field_specs(fresh_state)
    process.field_specs["Material name"].value = "PCBM"
    process.field_specs["Rotation speed [rpm]"].value = 1500
    process.field_specs["Rotation speed 1 [rpm]"] = dm.ProcessFieldSpec(
        key="Rotation speed 1 [rpm]"
    )

    filled, total = dm.compute_process_progress(process)
    field_keys = list(process.field_specs.keys())
    units_without_merging = len(field_keys)

    assert total < units_without_merging  # rotation speed/time/acceleration each merged


def test_compute_process_progress_excludes_not_required_fields(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    filled_before, total_before = dm.compute_process_progress(process)

    dm.set_field_required_for_progress(process.field_specs["Recipe file"], False)

    filled_after, total_after = dm.compute_process_progress(process)
    assert total_after == total_before - 1
    assert filled_after <= filled_before


def test_set_field_required_for_progress_does_not_touch_value_or_provenance(fresh_state):
    spec = dm.ProcessFieldSpec(key="Notes", value="something", is_outlier=True)
    dm.set_field_required_for_progress(spec, False)
    assert spec.value == "something"
    assert spec.is_outlier is True
    assert spec.required_for_progress is False


def test_compute_experiment_progress_sums_across_processes(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 0, "solutes": 0})
    fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.get_process(1).field_specs["Material name"].value = "PCBM"

    filled, total = dm.compute_experiment_progress(fresh_state)
    process_1_filled, process_1_total = dm.compute_process_progress(fresh_state.get_process(1))
    process_2_filled, process_2_total = dm.compute_process_progress(fresh_state.get_process(2))

    assert filled == process_1_filled + process_2_filled
    assert total == process_1_total + process_2_total


def test_compute_experiment_info_progress_excludes_only_computed_keys(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    fresh_state.experiment_info_fields["Project_Name"].value = "CsFA"
    # isolate this test from config/required_fields.json's "Notes" exception - it's
    # testing the computed-key exclusion specifically, not the required-fields config
    fresh_state.experiment_info_fields["Notes"].required_for_progress = True

    filled, total = dm.compute_experiment_info_progress(fresh_state)

    all_keys = set(fresh_state.experiment_info_fields)
    relevant_keys = all_keys - dm.EXPERIMENT_INFO_COMPUTED_KEYS
    assert total == len(relevant_keys)
    assert filled == 1


def test_compute_experiment_info_progress_excludes_not_required_fields(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    total_before = dm.compute_experiment_info_progress(fresh_state)[1]

    dm.set_field_required_for_progress(fresh_state.experiment_info_fields["Project_Name"], False)

    total_after = dm.compute_experiment_info_progress(fresh_state)[1]
    assert total_after == total_before - 1


# ---------------------------------------------------------------------------
# compute_sample_set_split -- "most natural division" preloaded in SampleSetupPanel's
# per-set count inputs.
# ---------------------------------------------------------------------------


def test_compute_sample_set_split_even():
    assert dm.compute_sample_set_split(12, 3) == [4, 4, 4]


def test_compute_sample_set_split_uneven_matches_spec_example():
    assert dm.compute_sample_set_split(15, 4) == [4, 4, 4, 3]


def test_compute_sample_set_split_zero_sets_returns_empty():
    assert dm.compute_sample_set_split(10, 0) == []


# ---------------------------------------------------------------------------
# progress_band -- color-coding thresholds for ProgressBarWidget.
# ---------------------------------------------------------------------------


def test_progress_band_empty_denominator_is_red():
    assert dm.progress_band(0, 0) == "red"


def test_progress_band_thresholds():
    assert dm.progress_band(1, 3) == "red"  # exactly 1/3
    assert dm.progress_band(34, 100) == "yellow"  # just over 1/3
    assert dm.progress_band(66, 100) == "yellow"  # just under 2/3
    assert dm.progress_band(67, 100) == "blue"  # just over 2/3
    assert dm.progress_band(89, 100) == "blue"
    assert dm.progress_band(90, 100) == "green"


# ---------------------------------------------------------------------------
# resolve_process_type -- the real NOMAD archive's 'method' string doesn't always match
# this app's AVAILABLE_PROCESSES labels 1:1 (the method each nomad-baseclasses class sets,
# e.g. "Atomic Layer Deposition" and a single "Cleaning"), and some process classes have
# no method at all and are recognised by their m_def instead.
# ---------------------------------------------------------------------------


def test_resolve_process_type_passthrough_for_exact_match():
    assert dm.resolve_process_type({"method": "Spin Coating"}) == "Spin Coating"
    assert (
        dm.resolve_process_type({"method": "Close Space Sublimation"}) == "Close Space Sublimation"
    )


def test_resolve_process_type_aliases_inkjet_printing():
    assert dm.resolve_process_type({"method": "Inkjet printing"}) == "Inkjet Printing"


def test_resolve_process_type_evaporation_vs_co_evaporation():
    plain = {"method": "Evaporation", "co_evaporation": False, "organic_evaporation": [{}]}
    co_flag = {"method": "Evaporation", "co_evaporation": True}
    co_data = {"method": "Evaporation", "perovskite_evaporation": [{"thickness": 20.0}]}
    assert dm.resolve_process_type(plain) == "Evaporation"
    assert dm.resolve_process_type(co_flag) == "Co-Evaporation"
    assert dm.resolve_process_type(co_data) == "Co-Evaporation"


def _m_def(class_name: str) -> str:
    """An m_def the way NOMAD writes it: the schema package path plus the class name."""
    return f"nomad_perotf.schema_packages.perotf_package.{class_name}"


def test_resolve_process_type_m_def_maps_the_three_method_less_classes():
    from perotf_utils.config import ENTRY_TYPES

    cases = {
        ENTRY_TYPES["process"]: "Generic Process",
        ENTRY_TYPES["thermal_annealing"]: "Annealing",
        ENTRY_TYPES["lamination"]: "Lamination",
    }
    for class_name, expected in cases.items():
        assert dm.resolve_process_type({"m_def": _m_def(class_name)}) == expected, class_name
        assert expected in dm.AVAILABLE_PROCESSES


def test_resolve_process_type_m_def_matches_lab_variants_of_thermal_annealing():
    """The lab variants (e.g. CR_ and TFL_) of the thermal annealing class share its
    suffix and have no method either."""
    from perotf_utils.config import ENTRY_TYPES

    prefix, suffix = ENTRY_TYPES["thermal_annealing"].split("_", 1)
    for lab in ("CR", "TFL"):
        step = {"method": None, "m_def": _m_def(f"{prefix}_{lab}_{suffix}")}
        assert dm.resolve_process_type(step) == "Annealing", lab


def test_resolve_process_type_m_def_unknown_class_returns_none():
    assert dm.resolve_process_type({"m_def": _m_def("SomethingElse")}) is None


def test_resolve_process_type_aliases_atomic_layer_deposition_to_ald():
    assert dm.resolve_process_type({"method": "Atomic Layer Deposition"}) == "ALD"


def test_resolve_process_type_unknown_method_returns_none():
    assert dm.resolve_process_type({"method": "Some Future Process"}) is None
    assert dm.resolve_process_type({}) is None


def test_resolve_process_type_cleaning_prefers_uv_ozone_when_uv_has_data():
    step = {"method": "Cleaning", "cleaning_uv": [{"time": 15.0}], "cleaning_plasma": [{}]}
    assert dm.resolve_process_type(step) == "Cleaning UV-Ozone"


def test_resolve_process_type_cleaning_uses_o2_plasma_when_only_plasma_has_data():
    step = {
        "method": "Cleaning",
        "cleaning_uv": [{}],
        "cleaning_plasma": [{"gas": "Oxygen", "time": 180.0}],
    }
    assert dm.resolve_process_type(step) == "Cleaning O2-Plasma"


def test_resolve_process_type_cleaning_defaults_to_uv_ozone_when_neither_has_data():
    step = {"method": "Cleaning", "cleaning_uv": [{}], "cleaning_plasma": [{}]}
    assert dm.resolve_process_type(step) == "Cleaning UV-Ozone"


ALD_STEP = {
    "method": "Atomic Layer Deposition",
    "positon_in_experimental_plan": 8.0,
    "location": "HyALD",
    "layer": [{"layer_type": "ETL-buffer layer", "layer_material_name": "SnOx"}],
    "properties": {
        "source": "SnOx",
        "thickness": 20.0,
        "temperature": 80.0,
        "rate": 0.05,
        "time": 1800.0,
        "number_of_cycles": 140,
        "material": {
            "pulse_duration": 1.0,
            "manifold_temperature": 80.0,
            "bottle_temperature": 60.0,
            "material": {"name": "TDMASn"},
        },
        "oxidizer_reducer": {
            "pulse_duration": 0.2,
            "manifold_temperature": 80.0,
            "material": {"name": "H2O"},
        },
    },
}

SPUTTERING_STEP = {
    "method": "Sputtering",
    "positon_in_experimental_plan": 9.0,
    "location": "Sputter tool",
    "layer": [{"layer_type": "Electron Transport Layer", "layer_material_name": "TiO2"}],
    "processes": [
        {
            "thickness": 50.0,
            "pressure": 0.01,
            "temperature": 200.0,
            "burn_in_time": 60.0,
            "deposition_time": 300.0,
            "power": 150.0,
            "gas_flow_rate": 20.0,
            "rotation_rate": 30.0,
            "gas_2": {"name": "Argon"},
        }
    ],
}

CLEANING_STEP = {
    "method": "Cleaning",
    "positon_in_experimental_plan": 1.0,
    "cleaning": [
        {"time": 10.0, "name": "Hellmanex-DI water", "temperature": 30.0},
        {"time": 10.0, "name": "DI Water", "temperature": 30.0},
    ],
    "cleaning_uv": [{"time": 15.0}],
    "cleaning_plasma": [{}],
}


def test_fetch_process_field_values_ald_step():
    cache = _cache_with("B1", [ALD_STEP])
    values, _source = dm.fetch_process_field_values("url", "token", cache, "B1", "ALD")
    assert values["Material name"] == "SnOx"
    assert values["Source"] == "SnOx"
    assert values["Number of cycles"] == 140
    assert values["Precursor 1"] == "TDMASn"
    assert values["Precursor 2 (Oxidizer/Reducer)"] == "H2O"


def test_fetch_process_field_values_sputtering_step():
    cache = _cache_with("B1", [SPUTTERING_STEP])
    values, _source = dm.fetch_process_field_values("url", "token", cache, "B1", "Sputtering")
    assert values["Material name"] == "TiO2"
    assert values["Gas"] == "Argon"
    assert values["Power [W]"] == 150.0


def test_fetch_process_field_values_cleaning_uv_ozone_step():
    """CleaningTechnique.time is declared unit='minute' in nomad-baseclasses even though
    this app's Excel columns for it are labeled seconds - fetch_process_field_values
    applies the confirmed x60 conversion (process_specs.py's "multiply": 60)."""
    cache = _cache_with("B1", [CLEANING_STEP])
    values, _source = dm.fetch_process_field_values(
        "url", "token", cache, "B1", "Cleaning UV-Ozone"
    )
    assert values["Solvent 1"] == "Hellmanex-DI water"
    assert values["Solvent 2"] == "DI Water"
    assert values["Time 1 [s]"] == 600.0
    assert values["UV-Ozone Time [s]"] == 900.0


def test_build_process_sequence_from_batch_resolves_aliased_and_cleaning_types(fresh_state):
    cache = _cache_with("B1", [CLEANING_STEP, ALD_STEP, SPUTTERING_STEP])
    sequence = dm.build_process_sequence_from_batch("url", "token", cache, "B1")
    assert [p.process_type for p in sequence] == ["Cleaning UV-Ozone", "ALD", "Sputtering"]
    # Cleaning's config was widened to match its 2 real solvent entries
    assert sequence[0].config["solvents"] == 2


# ---------------------------------------------------------------------------
# Process-type mappings sourced from the nomad-baseclasses map_<type> functions the
# experiment parser calls, not sample archive dumps - see process_specs.py's docstring.
# ---------------------------------------------------------------------------


def test_fetch_process_field_values_evaporation_falls_back_across_organic_inorganic():
    """Regression test for a real gap: Evaporation's field values live under
    organic_evaporation OR inorganic_evaporation depending on an 'Organic' flag this app
    doesn't itself track - same field names inside, different parent key. Before the
    'paths' (plural) fallback, only organic_evaporation was ever checked, so any
    inorganic evaporation (e.g. a metal electrode) silently returned nothing."""
    organic_step = {
        "method": "Evaporation",
        "layer": [{"layer_material_name": "PCBM", "layer_type": "ETL"}],
        "organic_evaporation": [{"thickness": 100.0, "start_rate": 0.5}],
    }
    inorganic_step = {
        "method": "Evaporation",
        "layer": [{"layer_material_name": "Gold", "layer_type": "Electrode"}],
        "inorganic_evaporation": [{"thickness": 80.0, "start_rate": 1.2}],
    }
    organic_cache = _cache_with("B1", [organic_step])
    inorganic_cache = _cache_with("B2", [inorganic_step])

    organic_values, _ = dm.fetch_process_field_values(
        "url", "token", organic_cache, "B1", "Evaporation"
    )
    inorganic_values, _ = dm.fetch_process_field_values(
        "url", "token", inorganic_cache, "B2", "Evaporation"
    )

    assert organic_values["Thickness [nm]"] == 100.0
    assert inorganic_values["Thickness [nm]"] == 80.0
    assert inorganic_values["Material name"] == "Gold"


CO_EVAPORATION_STEP = {
    "method": "Co-Evaporation",
    "layer": [{"layer_material_name": "Aluminium", "layer_type": "Electrode"}],
    "perovskite_evaporation": [
        {"chemical_2": {"name": "Copper"}, "thickness": 20.0, "target_rate": 1.5},
        {"chemical_2": {"name": "Silver"}, "thickness": 21.0, "target_rate": 1.6},
    ],
}


def test_fetch_process_field_values_co_evaporation_indexed_materials():
    cache = _cache_with("B1", [CO_EVAPORATION_STEP])
    values, _source = dm.fetch_process_field_values("url", "token", cache, "B1", "Co-Evaporation")
    assert values["Material name"] == "Aluminium"
    assert values["Material name 1"] == "Copper"
    assert values["Thickness 1 [nm]"] == 20.0
    assert values["Material name 2"] == "Silver"
    assert values["Rate 2 [angstrom/s]"] == 1.6


SLOT_DIE_COATING_STEP = {
    "method": "Slot Die Coating",
    "layer": [{"layer_material_name": "Perovskite", "layer_type": "Absorber"}],
    "solution": [
        {
            "solution_details": {
                "solvent": [{"name": "DMF 0.01 milliliter", "chemical_2": {"name": "DMF"}}]
            }
        }
    ],
    "properties": {
        "flow_rate": 25.0,
        "slot_die_head_speed": 15.0,
        "coating_run": "R1",
        "coated_area": 100.0,
    },
    "quenching": {
        "air_knife_angle": 45.0,
        "drying_gas_temperature": 25.0,
        "heat_transfer_coefficient": 10.0,
    },
}


def test_fetch_process_field_values_slot_die_coating():
    cache = _cache_with("B1", [SLOT_DIE_COATING_STEP])
    values, _source = dm.fetch_process_field_values("url", "token", cache, "B1", "Slot Die Coating")
    assert values["Flow rate [uL/min]"] == 25.0
    assert values["Speed [mm/s]"] == 15.0
    assert values["Solvent 1 name"] == "DMF"
    assert values["Coating run"] == "R1"
    assert values["Coated area [mm²]"] == 100.0
    assert values["Air knife angle [°]"] == 45.0
    assert values["Drying gas temperature [°]"] == 25.0
    assert values["Heat transfer coefficient [W m^-2 K^-1]"] == 10.0


DIP_COATING_STEP = {
    "method": "Dip Coating",
    "layer": [{"layer_material_name": "Perovskite", "layer_type": "Absorber"}],
    "solution": [
        {
            "solution_details": {
                "solvent": [
                    {"name": "DMF 0.05", "chemical_2": {"name": "DMF"}, "chemical_volume": 0.05}
                ],
                "solute": [{"chemical_2": {"name": "PbI2"}, "concentration_mol": 0.001}],
            }
        }
    ],
    "properties": {"time": 15.0},
}


def test_fetch_process_field_values_dip_coating():
    cache = _cache_with("B1", [DIP_COATING_STEP])
    values, _source = dm.fetch_process_field_values("url", "token", cache, "B1", "Dip Coating")
    assert values["Dipping duration [s]"] == 15.0
    assert values["Solvent 1 name"] == "DMF"
    assert values["Solvent 1 volume [uL]"] == 0.05
    assert values["Solute 1 type"] == "PbI2"


INKJET_PRINTING_STEP = {
    "method": "Inkjet printing",
    "layer": [{"layer_material_name": "PEDOT:PSS", "layer_type": "HTL"}],
    "properties": {
        "printing_run": "Run 2",
        "cartridge_pressure": 300.0,
        "print_head_properties": {"print_head_name": "Spectra 0.8uL"},
    },
    "print_head_path": {"quality_factor": 3},
    "atmosphere": {"temperature": 21.0, "relative_humidity": 35.0},
}


def test_fetch_process_field_values_inkjet_printing():
    cache = _cache_with("B1", [INKJET_PRINTING_STEP])
    values, _source = dm.fetch_process_field_values("url", "token", cache, "B1", "Inkjet Printing")
    assert values["Printhead name"] == "Spectra 0.8uL"
    assert values["Printing run"] == "Run 2"
    assert values["Ink reservoir pressure [mbar]"] == 300.0
    assert values["Quality factor"] == 3
    assert values["Room temperature [°C]"] == 21.0
    assert values["rel. humidity [%]"] == 35.0


def test_inkjet_printing_quenching_blocks_have_no_archive_paths():
    """The parser stores no gas/vacuum quenching for inkjet printing, so its checkboxes
    add columns but never autofill them."""
    paths = dm.PROCESS_TYPE_FIELD_PATHS["Inkjet Printing"]
    assert "Gas" not in paths
    assert "Vacuum quenching start time [s]" not in paths


CSS_STEP = {
    "method": "Close Space Sublimation",
    "positon_in_experimental_plan": 4.0,
    "location": "CSS tool",
    "layer": [{"layer_material_name": "CdTe", "layer_type": "Absorber"}],
    "process": {
        "pressure": 40.0,
        "source_temperature": 600.0,
        "substrate_temperature": 500.0,
        "substrate_source_distance": 4.0,
        "thickness": 500.0,
        "deposition_time": 120.0,
        "carrier_gas": "no",
        "material_state": "Powder",
        "source_material_mixture": [
            {"source_material": {"name": "CdTe"}, "mixing_ratio": 3.0},
            {"source_material": {"name": "CdSe"}, "mixing_ratio": 1.0},
        ],
        "process_preparation": {"rotation_speed": 300.0, "rotation_time": 30.0},
    },
}


def test_fetch_process_field_values_close_space_sublimation():
    cache = _cache_with("B1", [CSS_STEP])
    values, _source = dm.fetch_process_field_values(
        "url", "token", cache, "B1", "Close Space Sublimation"
    )
    assert values["Material name"] == "CdTe"
    assert values["Process pressure [mbar]"] == 40.0
    assert values["Source temperature [°C]"] == 600.0
    assert values["Substrate source distance [mm]"] == 4.0
    assert values["Carrier gas"] == "no"
    assert values["Material state"] == "Powder"
    assert values["Source material name 1"] == "CdTe"
    assert values["Source material name 2"] == "CdSe"
    assert values["Source material ratio 2"] == 1.0
    assert values["Milling rotation speed [rpm]"] == 300.0


def test_infer_config_from_source_step_css_materials_and_milling():
    inferred = dm.infer_config_from_source_step("Close Space Sublimation", CSS_STEP)
    assert inferred == {"materials": 2, "milling": True}


LAMINATION_STEP = {
    "location": "Laminator",
    "positon_in_experimental_plan": 9.0,
    "settings": {"pressure": 1.5, "force": 100.0, "heat_up_time": 60.0},
}


def test_fetch_process_field_values_lamination():
    from perotf_utils.config import ENTRY_TYPES

    step = dict(LAMINATION_STEP, m_def=f"pkg.{ENTRY_TYPES['lamination']}")
    cache = _cache_with("B1", [step])
    values, _source = dm.fetch_process_field_values("url", "token", cache, "B1", "Lamination")
    assert values == {
        "Tool/GB name": "Laminator",
        "Pressure [MPa]": 1.5,
        "Force [N]": 100.0,
        "Heat up time [s]": 60.0,
    }


LASER_SCRIBING_STEP = {
    "method": "Laser Scribing",
    "recipe_file": "test_scribing_recipe.xml",
    "properties": {"laser_wavelength": 532, "speed": 100.0},
}


def test_fetch_process_field_values_laser_scribing():
    cache = _cache_with("B1", [LASER_SCRIBING_STEP])
    values, _source = dm.fetch_process_field_values("url", "token", cache, "B1", "Laser Scribing")
    assert values["Recipe file"] == "test_scribing_recipe.xml"
    assert values["Laser wavelength [nm]"] == 532
    assert values["Speed [mm/s]"] == 100.0


ANNEALING_STEP = {
    "method": "Annealing",
    "annealing": {"temperature": 150.0, "atmosphere": "Nitrogen"},
    "atmosphere": {"relative_humidity": 35.0},
}


def test_fetch_process_field_values_annealing():
    cache = _cache_with("B1", [ANNEALING_STEP])
    values, _source = dm.fetch_process_field_values("url", "token", cache, "B1", "Annealing")
    assert values["Annealing temperature [°C]"] == 150.0
    assert values["Annealing athmosphere"] == "Nitrogen"
    assert values["Relative humidity [%]"] == 35.0


def test_fetch_process_field_values_cleaning_o2_plasma_gas_plasma_fields():
    step = {
        "method": "Cleaning",
        "cleaning_uv": [{}],
        "cleaning_plasma": [{"plasma_type": "Oxygen", "time": 180.0, "power": 50.0}],
    }
    cache = _cache_with("B1", [step])
    values, _source = dm.fetch_process_field_values(
        "url", "token", cache, "B1", "Cleaning O2-Plasma"
    )
    assert values["Gas-Plasma Gas"] == "Oxygen"
    # raw 180.0 is minutes (CleaningTechnique.time unit) -> x60 to the "[s]" column
    assert values["Gas-Plasma Time [s]"] == 10800.0
    assert values["Gas-Plasma Power [W]"] == 50.0


# ---------------------------------------------------------------------------
# GUI smoke tests -- light coverage only (matches this repo's precedent for widget
# code); the heavy logic is already covered on the data_manager layer above. These focus
# on the interactions that live partly in gui_components.py itself (event wiring).
# ---------------------------------------------------------------------------


def test_process_fields_panel_renders_one_row_per_field(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    panel = gc.ProcessFieldsPanel(fresh_state, process)
    # children[0] is the one-line provenance summary, ahead of one row per field
    assert len(panel.children) == len(process.field_specs) + 1


def test_process_fields_panel_varies_checkbox_promotes_scope(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    process.field_specs["Material name"].value = "C60"
    panel = gc.ProcessFieldsPanel(fresh_state, process)

    panel._on_varies_change("Material name", True)

    spec = process.field_specs["Material name"]
    assert spec.varies is True
    assert spec.per_sample_values == {1: "C60"}  # seeded from prior constant value


def test_process_fields_panel_value_edit_sets_manual_provenance(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    panel = gc.ProcessFieldsPanel(fresh_state, process)

    panel._on_value_change("Material name", "hand-typed")

    spec = process.field_specs["Material name"]
    assert spec.value == "hand-typed"
    assert spec.provenance.source == "manual"


def test_process_fields_panel_calls_on_change_callback(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    calls = []
    panel = gc.ProcessFieldsPanel(fresh_state, process, on_change=lambda: calls.append(1))

    panel._on_value_change("Material name", "PCBM")

    assert calls == [1]


def test_process_fields_panel_required_field_shows_asterisk(fresh_state):
    """The per-row 'Required' checkbox was removed (required_for_progress is now
    config/required_fields.json-driven, not user-toggleable in this panel) - a required
    field's label instead gets a trailing '*', an excepted one doesn't."""
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    dm.set_field_required_for_progress(process.field_specs["Recipe file"], False)

    panel = gc.ProcessFieldsPanel(fresh_state, process)

    # panel.children[0] is the provenance summary, not a field row - skip it
    labels = [row.children[1].value for row in panel.children[1:]]
    assert "Recipe file" in labels  # excepted -> no asterisk
    assert "Laser wavelength [nm] *" in labels  # still required by default -> asterisk


def test_process_fields_panel_shows_preview_placeholder_for_empty_field(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    fresh_state.whole_experiment_template_batch_id = "B1"
    cache = _cache_with("B1", [SPIN_COATING_STEP])

    panel = gc.ProcessFieldsPanel(fresh_state, process, cache=cache)

    rows_by_label = {row.children[1].value: row for row in panel.children[1:]}
    material_value_widget = rows_by_label["Material name *"].children[2]
    assert material_value_widget.value == ""
    assert material_value_widget.placeholder == "Me4PACz"


def test_process_fields_panel_falls_back_to_generic_placeholder_without_preview(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)

    panel = gc.ProcessFieldsPanel(fresh_state, process, cache=dm.NomadSessionCache())

    rows_by_label = {row.children[1].value: row for row in panel.children[1:]}
    material_value_widget = rows_by_label["Material name *"].children[2]
    assert material_value_widget.placeholder == "value"


def test_varying_fields_matrix_shows_placeholder_when_nothing_varies(fresh_state):
    matrix = gc.VaryingFieldsMatrix(fresh_state)
    assert len(matrix.children) == 1  # placeholder HTML only


def test_varying_fields_matrix_renders_header_and_sample_rows(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    fresh_state.add_sample(variation_group_index=0, sample_number=2)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())

    matrix = gc.VaryingFieldsMatrix(fresh_state)
    field_label = dm.iter_varying_fields(fresh_state)[0][0]

    assert matrix.header_widget(field_label) is not None
    assert matrix.cell_widget(1, field_label) is not None
    assert matrix.cell_widget(2, field_label) is not None


def _add_manual_field(process, key):
    """peroTF's Excel_creator writes no Datetime/Operator columns, but the quick-fill
    buttons are keyed by field name; these tests add such a field by hand."""
    spec = dm.ProcessFieldSpec(key=key)
    process.field_specs[key] = spec
    return spec


def test_varying_fields_matrix_operator_cell_gets_a_me_button(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    operator = _add_manual_field(process, "Operator")
    dm.set_field_varies(operator, True, fresh_state.sample_numbers())

    matrix = gc.VaryingFieldsMatrix(fresh_state)
    field_label = dm.iter_varying_fields(fresh_state)[0][0]
    operator_cell_slot = matrix.cell_widget(1, field_label)

    text_widget, quick_fill_button = operator_cell_slot.children
    assert quick_fill_button.description == "Me"

    quick_fill_button.click()
    assert text_widget.value  # filled in by the click


def test_varying_fields_matrix_column_header_has_populate_button(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    fresh_state.add_sample(variation_group_index=0, sample_number=2)
    spec = process.field_specs["Material name"]
    dm.set_field_varies(spec, True, fresh_state.sample_numbers())

    matrix = gc.VaryingFieldsMatrix(fresh_state)
    field_label = dm.iter_varying_fields(fresh_state)[0][0]
    material_header = matrix.header_widget(field_label)
    _label, populate_button = material_header.children
    assert populate_button.icon == "arrow-down"
    assert populate_button.tabbable is False  # excluded from the column-major Tab cycle

    dm.set_field_manual(spec, "C60", sample_number=1)
    populate_button.click()

    assert spec.per_sample_values == {1: "C60", 2: "C60"}


def test_varying_fields_matrix_shows_set_column_before_variation(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=2, sample_number=1)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())

    matrix = gc.VaryingFieldsMatrix(fresh_state)

    columns = matrix.columns()
    assert columns.index("Subbatch") < columns.index("Variation")
    # header_widget returns the uniform [text, button-or-placeholder] VBox every column
    # header now shares (see _header_for's docstring) - the text block is children[0].
    assert "Subbatch" in matrix.header_widget("Subbatch").children[0].value
    # 1-based: variation_group_index=2 displays as Subbatch 3, matching subbatch_for_sample
    assert matrix.cell_widget(1, "Subbatch").value == "3"


def test_varying_fields_matrix_tab_order_is_column_major(fresh_state):
    """Product ask: pressing Tab should move to the next SAMPLE within the same field,
    not the next field for the same sample - achieved by inserting cells into the
    GridBox in column-major order (visual position is pinned separately via
    grid_row/grid_column, see the class docstring). Doesn't assume which field's column
    comes first (that's just field_specs' own dict order) - only that, WITHIN one
    field's column, consecutive samples are inserted back-to-back with nothing from any
    other column interleaved."""
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    fresh_state.add_sample(variation_group_index=0, sample_number=2)
    fresh_state.add_sample(variation_group_index=0, sample_number=3)
    spec = process.field_specs["Material name"]
    dm.set_field_varies(spec, True, fresh_state.sample_numbers())
    dm.set_field_varies(process.field_specs["Layer type"], True, fresh_state.sample_numbers())

    matrix = gc.VaryingFieldsMatrix(fresh_state)
    labels_by_spec_id = {id(s): label for label, s in dm.iter_varying_fields(fresh_state)}
    material_label = labels_by_spec_id[id(spec)]

    order = matrix.tab_order()
    material_indices = [
        order.index(matrix.cell_widget(sample_number, material_label))
        for sample_number in (1, 2, 3)
    ]

    assert material_indices == sorted(material_indices)
    assert material_indices[-1] - material_indices[0] == 2  # no other column interleaved


def test_varying_fields_matrix_hard_refresh_clears_before_rebuilding(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())
    matrix = gc.VaryingFieldsMatrix(fresh_state)
    field_label = dm.iter_varying_fields(fresh_state)[0][0]
    original_cell = matrix.cell_widget(1, field_label)

    matrix.hard_refresh()

    # children were cleared and rebuilt with fresh widgets, not left stale or duplicated
    assert matrix.cell_widget(1, field_label) is not original_cell


def test_varying_fields_matrix_hard_refresh_reflects_state_changes(fresh_state):
    matrix = gc.VaryingFieldsMatrix(fresh_state)
    assert len(matrix.children) == 1  # placeholder only, nothing varies yet

    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())

    matrix.hard_refresh()

    field_label = dm.iter_varying_fields(fresh_state)[0][0]
    assert matrix.cell_widget(1, field_label) is not None


def test_varying_fields_matrix_cell_edit_updates_state_without_touching_variation(fresh_state):
    # Variation is no longer auto-recomputed on every matrix edit (see
    # update_variation_column's docstring) - a plain field-cell edit only updates that
    # field's own per-sample value, nothing else.
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    spec = process.field_specs["Material name"]
    dm.set_field_varies(spec, True, fresh_state.sample_numbers())

    matrix = gc.VaryingFieldsMatrix(fresh_state)
    matrix._on_cell_change(spec, 1, "C60")

    assert spec.per_sample_values[1] == "C60"
    assert fresh_state.experiment_info_fields["Variation"].per_sample_values == {}


def test_varying_fields_matrix_variation_cell_manual_edit_is_respected(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())

    matrix = gc.VaryingFieldsMatrix(fresh_state)
    matrix._on_variation_cell_change(1, "custom label")

    assert fresh_state.experiment_info_fields["Variation"].per_sample_values[1] == "custom label"
    # subsequent field edit must not clobber the manually set Variation
    matrix._on_cell_change(process.field_specs["Material name"], 1, "C60")
    assert fresh_state.experiment_info_fields["Variation"].per_sample_values[1] == "custom label"


def test_auto_fill_variation_column_recomputes_after_later_field_edits(fresh_state):
    # Regression test for the reported bug: typing into a second varying column used to
    # do nothing to Variation once it had already been computed once (no-clobber blocked
    # every later recompute). auto_fill_variation_column is the new, explicit, full
    # recompute - it must pick up a field edited AFTER the first auto-fill too.
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    material = process.field_specs["Material name"]
    solvent = process.field_specs["Solvent 1 name"]
    dm.set_field_varies(material, True, fresh_state.sample_numbers())
    dm.set_field_varies(solvent, True, fresh_state.sample_numbers())

    dm.set_field_manual(material, "C60", sample_number=1)
    first_count = dm.auto_fill_variation_column(fresh_state)
    first_label = fresh_state.experiment_info_fields["Variation"].per_sample_values[1]

    dm.set_field_manual(solvent, "DMF", sample_number=1)
    second_count = dm.auto_fill_variation_column(fresh_state)
    second_label = fresh_state.experiment_info_fields["Variation"].per_sample_values[1]

    assert first_count == 1
    assert second_count == 1
    assert first_label != second_label
    assert "dmf" in second_label.lower()


def test_populate_column_from_first_copies_first_sample_into_the_rest(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"spinsteps": 1})
    dm.rebuild_field_specs(fresh_state)
    spec = process.field_specs["Rotation speed [rpm]"]
    for sample_number in (1, 2, 3):
        fresh_state.add_sample(variation_group_index=0, sample_number=sample_number)
    dm.set_field_varies(spec, True, fresh_state.sample_numbers())
    dm.set_field_manual(spec, "1000", sample_number=1)

    written = dm.populate_column_from_first(spec, fresh_state.sample_numbers())

    assert written == 2
    assert spec.per_sample_values == {1: "1000", 2: "1000", 3: "1000"}


def test_populate_column_from_first_noop_when_first_sample_empty(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"spinsteps": 1})
    dm.rebuild_field_specs(fresh_state)
    spec = process.field_specs["Rotation speed [rpm]"]
    for sample_number in (1, 2):
        fresh_state.add_sample(variation_group_index=0, sample_number=sample_number)
    dm.set_field_varies(spec, True, fresh_state.sample_numbers())

    written = dm.populate_column_from_first(spec, fresh_state.sample_numbers())

    assert written == 0
    assert spec.per_sample_values.get(2) is None


def test_fill_all_date_and_operator_fields_fills_experiment_info_and_processes(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    _add_manual_field(process, "Datetime")
    _add_manual_field(process, "Operator")

    date_count, operator_count = dm.fill_all_date_and_operator_fields(fresh_state, "Jane Doe")

    date_value = fresh_state.experiment_info_fields["Date"].value
    assert len(date_value) == 8 and date_value.isdigit()  # YYYYMMDD, a Nomad ID segment
    assert process.field_specs["Datetime"].value
    assert process.field_specs["Operator"].value == "Jane Doe"
    assert date_count >= 2  # Experiment Info's Date + the process's Datetime
    assert operator_count >= 1


def test_fill_all_date_and_operator_fields_fills_every_varying_sample(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    fresh_state.add_sample(variation_group_index=0, sample_number=2)
    spec = _add_manual_field(process, "Operator")
    dm.set_field_varies(spec, True, fresh_state.sample_numbers())

    dm.fill_all_date_and_operator_fields(fresh_state, "Jane Doe")

    assert spec.per_sample_values == {1: "Jane Doe", 2: "Jane Doe"}


def test_fill_all_date_and_operator_fields_overwrites_existing_values(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    fresh_state.experiment_info_fields["Date"].value = "20200101"

    dm.fill_all_date_and_operator_fields(fresh_state, "Jane Doe")

    assert fresh_state.experiment_info_fields["Date"].value != "20200101"


def test_missing_critical_fields_lists_all_when_empty(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    assert dm.missing_critical_fields(fresh_state) == ["Batch", "Project_Name", "Date"]


def test_missing_critical_fields_narrows_as_fields_are_filled(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    fresh_state.experiment_info_fields["Batch"].value = "1"
    assert dm.missing_critical_fields(fresh_state) == ["Project_Name", "Date"]

    fresh_state.experiment_info_fields["Project_Name"].value = "Test"
    assert dm.missing_critical_fields(fresh_state) == ["Date"]

    fresh_state.experiment_info_fields["Date"].value = "20260526"
    assert dm.missing_critical_fields(fresh_state) == []


def test_progress_bar_widget_reflects_compute_experiment_progress(fresh_state):
    fresh_state.add_process("Laser Scribing")  # not material-gated, always counts
    dm.rebuild_field_specs(fresh_state)
    bar = gc.ProgressBarWidget(fresh_state)

    filled, total = dm.compute_experiment_progress(fresh_state)
    assert bar.bar.value == filled
    assert bar.bar.max == total
    assert str(filled) in bar.label.value


def test_progress_bar_widget_refresh_updates_after_state_change(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    bar = gc.ProgressBarWidget(fresh_state)
    before = bar.bar.value

    process.field_specs["Recipe file"].value = "done"
    bar.refresh()

    assert bar.bar.value == before + 1


def test_progress_bar_widget_bar_style_and_message_reflect_band(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    bar = gc.ProgressBarWidget(fresh_state)
    assert bar.bar.bar_style == "danger"  # nothing filled yet -> red band
    assert bar.message.value

    for spec in process.field_specs.values():
        spec.value = "filled"
    bar.refresh()

    assert bar.bar.bar_style == "success"  # fully filled -> green band


def test_create_finish_section_places_progress_bar_above_buttons(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    cache = dm.NomadSessionCache()
    bar = gc.ProgressBarWidget(fresh_state)

    with patch.object(dm, "get_all_uploads", return_value=[]):
        section = gc.create_finish_section(fresh_state, "url", "token", cache, bar)

    (
        _skip_checkbox,
        _caption,
        _upload_dropdown,
        progress_bar_child,
        buttons_row,
        _nudge_area,
        _status_output,
        _download_js_output,
    ) = section.children
    assert progress_bar_child is bar
    assert len(buttons_row.children) == 3


def test_create_whole_experiment_template_picker_applies_selected_batch(fresh_state):
    cache = dm.NomadSessionCache()
    calls = []

    with (
        patch.object(dm, "get_batch_ids", return_value=["B1", "B2"]),
        patch.object(dm, "get_ids_in_batch", return_value=["S1"]),
        patch.object(dm, "get_processing_steps", return_value=[SPIN_COATING_STEP]),
    ):
        picker = gc.create_whole_experiment_template_picker(
            fresh_state, "url", "token", cache, on_change=lambda: calls.append(1)
        )
        # picker.children: [header, caption, batch_picker_vbox, progress_hbox];
        # batch_picker_vbox.children: [search_field, selector, load_button, status]
        batch_picker = picker.children[2]
        search_field, selector, load_button, status = batch_picker.children
        assert load_button.description == "Replicate Experiment"
        selector.value = ("B1",)
        load_button.click()

    assert fresh_state.get_process(1).field_specs["Material name"].value == "Me4PACz"
    assert fresh_state.whole_experiment_template_batch_id == "B1"
    assert calls == [1]
    assert "Replicated" in status.value
    assert "1 process(es)" in status.value

    progress_bar, progress_label = picker.children[3].children
    assert progress_bar.layout.visibility == "hidden"  # hidden again once done
    assert progress_label.value == ""


def test_create_whole_experiment_template_picker_shows_error_on_failure(fresh_state):
    cache = dm.NomadSessionCache()

    with (
        patch.object(dm, "get_batch_ids", return_value=["B1"]),
        patch.object(dm, "get_ids_in_batch", side_effect=RuntimeError("network down")),
    ):
        picker = gc.create_whole_experiment_template_picker(fresh_state, "url", "token", cache)
        batch_picker = picker.children[2]
        _search_field, selector, load_button, status = batch_picker.children
        selector.value = ("B1",)
        load_button.click()

    assert "Failed" in status.value
    assert "network down" in status.value


def test_process_sequence_builder_override_picker_shows_message_without_session(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    builder = gc.ProcessSequenceBuilder(fresh_state)  # no url/token/cache
    row = builder._process_rows[id(fresh_state.get_process(1))]
    picker = row._build_override_picker()
    assert "No NOMAD session" in picker.value


def test_process_sequence_builder_override_picker_applies_override_batch(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    cache = dm.NomadSessionCache()

    with (
        patch.object(dm, "get_batch_ids", return_value=["OVERRIDE_BATCH"]),
        patch.object(dm, "get_ids_in_batch", return_value=["S1"]),
        patch.object(dm, "get_processing_steps", return_value=[SPIN_COATING_STEP]),
    ):
        builder = gc.ProcessSequenceBuilder(fresh_state, "url", "token", cache)
        row = builder._process_rows[id(fresh_state.get_process(1))]
        batch_picker = row._build_override_picker()
        search_field, selector, load_button, _status = batch_picker.children
        selector.value = ("OVERRIDE_BATCH",)
        load_button.click()

    process = fresh_state.get_process(1)
    assert process.source_override_batch_id == "OVERRIDE_BATCH"
    assert process.field_specs["Material name"].value == "Me4PACz"


def test_process_sequence_builder_clear_override_button(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    dm.apply_process_override(
        fresh_state, process, "url", "token", _cache_with("B1", [SPIN_COATING_STEP]), "B1"
    )
    builder = gc.ProcessSequenceBuilder(fresh_state)

    row = builder._process_rows[id(process)]
    row._on_clear_override()

    assert process.source_override_batch_id is None


# ---------------------------------------------------------------------------
# app.py -- thin orchestrator smoke test
# ---------------------------------------------------------------------------


def test_initialize_ui_builds_widget_tree():
    with (
        patch.object(dm, "get_batch_ids", return_value=["B1"]),
        patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]),
    ):
        main_interface = app_module.initialize_ui("url", "token")
    assert len(main_interface.children) > 0


def test_initialize_ui_auto_applies_default_sample_setup():
    """Product ask: the Varying Fields table (and everything downstream) previously
    stayed empty until a user noticed and clicked 'Apply Sample Setup' themselves -
    initialize_ui now runs it once automatically with the panel's own defaults."""
    with (
        patch.object(dm, "get_batch_ids", return_value=["B1"]),
        patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]),
    ):
        main_interface = app_module.initialize_ui("url", "token")
    sample_setup = main_interface.children[2]
    assert len(sample_setup.state.samples) == sample_setup.total_samples_input.value > 0


def test_initialize_ui_closes_widgets_from_previous_call():
    with (
        patch.object(dm, "get_batch_ids", return_value=["B1"]),
        patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]),
    ):
        app_module.initialize_ui("url", "token")

        first_call_ids = set(app_module._ui_widget_ids)
        app_module.initialize_ui("url", "token")

    import ipywidgets as widgets

    for widget_id in first_call_ids:
        assert widget_id not in widgets.Widget.widgets


def test_initialize_ui_template_pick_refreshes_visible_sequence_rows():
    """End-to-end regression test for a real bug: ProcessSequenceBuilder only re-rendered
    itself in response to its own actions, so picking a whole-experiment template (which
    replaces state.process_sequence from outside that widget) updated the data but left
    the on-screen process rows stale - see app.py's refresh_all wiring."""
    with (
        patch.object(dm, "get_batch_ids", return_value=["B1"]),
        patch.object(dm, "get_ids_in_batch", return_value=["S1"]),
        patch.object(dm, "get_processing_steps", return_value=[SPIN_COATING_STEP]),
        patch.object(dm, "get_all_uploads", return_value=[]),
    ):
        main_interface = app_module.initialize_ui("url", "token")
        # app.py's main_interface.children order: h2, h4, sample_setup, template_picker,
        # quick_fill_all_button, h4, sequence_builder, ... (progress_bar now lives inside
        # finish_section, not at the top level)
        template_picker = main_interface.children[3]
        sequence_builder = main_interface.children[6]
        assert len(sequence_builder.rows_box.children) == 0

        batch_picker = template_picker.children[2]
        _search_field, selector, load_button, _status = batch_picker.children
        selector.value = ("B1",)
        load_button.click()

    assert len(sequence_builder.rows_box.children) == 1


def test_initialize_ui_refresh_table_button_updates_stale_matrix():
    """The 'Refresh Table' button is a manual fallback for when the Varying Fields
    matrix doesn't display for large datasets - clicking it must always pick up the
    current state, even if nothing else triggered a refresh in between."""
    with (
        patch.object(dm, "get_batch_ids", return_value=["B1"]),
        patch.object(dm, "get_all_uploads", return_value=[]),
    ):
        main_interface = app_module.initialize_ui("url", "token")
        # ... h4 Varying Fields, caption, refresh_matrix_button, refresh_matrix_status, matrix
        refresh_matrix_button = main_interface.children[9]
        matrix = main_interface.children[11]
        assert refresh_matrix_button.description == "Refresh Table"
        assert len(matrix.children) == 1  # placeholder, nothing varies yet

        # Simulate state changing without going through any widget that would normally
        # trigger matrix.refresh() itself. Sample Setup auto-applies its default sample
        # set on page load now, so state.samples is already populated here - no need to
        # add one manually.
        state = matrix.state
        sample_count = len(state.samples)
        assert sample_count > 0
        process = state.add_process("Evaporation")
        dm.rebuild_field_specs(state)
        dm.set_field_varies(process.field_specs["Material name"], True, state.sample_numbers())
        assert len(matrix.children) == 1  # not yet reflected

        refresh_matrix_button.click()

    field_label = dm.iter_varying_fields(state)[0][0]
    for sample_number in state.sample_numbers():
        assert matrix.cell_widget(sample_number, field_label) is not None


# ---------------------------------------------------------------------------
# Outlier flagging during autofill
# ---------------------------------------------------------------------------


def test_autofill_flags_outlier_far_from_distribution(fresh_state):
    """All 4 steps share one positon_in_experimental_plan (same sequence position, 4
    different samples) - occurrence 0 always autofills from whichever is FIRST in the
    list (distinct_steps_for_process_type's representative pick), so the extreme value
    is listed first here specifically to be the one fetched+written, with the other 3
    providing its comparison distribution (100/130/102 - not outliers relative to each
    other, but 99999 clearly is)."""
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with(
        "B1",
        [
            SPIN_COATING_STEP_EXTREME_OUTLIER,
            SPIN_COATING_STEP,
            SPIN_COATING_STEP_VARIATION,
            SPIN_COATING_STEP_CLOSE,
        ],
    )

    dm.autofill_process_from_batch(process, "url", "token", cache, "B1", occurrence=0)

    assert process.field_specs["Annealing temperature [°C]"].value == 99999.0
    assert process.field_specs["Annealing temperature [°C]"].is_outlier is True


def test_autofill_does_not_flag_value_within_distribution(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with(
        "B1",
        [
            SPIN_COATING_STEP,
            SPIN_COATING_STEP_VARIATION,
            SPIN_COATING_STEP_CLOSE,
            SPIN_COATING_STEP_EXTREME_OUTLIER,
        ],
    )

    dm.autofill_process_from_batch(process, "url", "token", cache, "B1", occurrence=0)

    assert process.field_specs["Annealing temperature [°C]"].value == 100.0
    assert process.field_specs["Annealing temperature [°C]"].is_outlier is False


def test_autofill_outlier_flag_cleared_by_manual_edit(fresh_state):
    # Extreme value listed first so occurrence 0 (the shared position's representative,
    # see distinct_steps_for_process_type) fetches+writes it, with the other two as its
    # comparison distribution.
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with(
        "B1",
        [SPIN_COATING_STEP_EXTREME_OUTLIER, SPIN_COATING_STEP, SPIN_COATING_STEP_CLOSE],
    )
    dm.autofill_process_from_batch(process, "url", "token", cache, "B1", occurrence=0)
    assert process.field_specs["Annealing temperature [°C]"].is_outlier is True

    dm.set_field_manual(process.field_specs["Annealing temperature [°C]"], 105)

    assert process.field_specs["Annealing temperature [°C]"].is_outlier is False


# ---------------------------------------------------------------------------
# build_nudge_queue / build_missing_fields_summary
# ---------------------------------------------------------------------------


def test_build_nudge_queue_prioritizes_worst_gap_process_first(fresh_state):
    small_gap = fresh_state.add_process("Annealing")  # few fields
    big_gap = fresh_state.add_process("Slot Die Coating", config={"solvents": 3, "solutes": 2})
    dm.rebuild_field_specs(fresh_state)
    small_gap.field_specs["Notes"].value = ""  # leave everything else missing too
    big_gap.field_specs["Material name"].value = ""

    queue = dm.build_nudge_queue(fresh_state)
    missing_items = [item for item in queue if item.kind == "missing"]

    # first item in the queue must belong to the process with more missing fields
    assert missing_items[0].sequence_index == big_gap.sequence_index


def test_build_nudge_queue_excludes_varying_fields(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())

    queue = dm.build_nudge_queue(fresh_state)

    assert not any(item.field_key == "Material name" for item in queue)


def test_build_nudge_queue_missing_items_come_before_outlier_items(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with(
        "B1", [SPIN_COATING_STEP_EXTREME_OUTLIER, SPIN_COATING_STEP, SPIN_COATING_STEP_CLOSE]
    )
    dm.autofill_process_from_batch(process, "url", "token", cache, "B1", occurrence=0)
    process.field_specs["Layer type"].value = ""  # still missing

    queue = dm.build_nudge_queue(fresh_state)
    kinds = [item.kind for item in queue]

    assert kinds.index("missing") < kinds.index("outlier")


def test_build_nudge_queue_respects_max_items(fresh_state):
    process = fresh_state.add_process("Slot Die Coating", config={"solvents": 3, "solutes": 2})
    dm.rebuild_field_specs(fresh_state)
    assert len(process.field_specs) > 2

    queue = dm.build_nudge_queue(fresh_state, max_items=2)

    assert len(queue) == 2


def test_build_nudge_queue_includes_material_gated_process_with_empty_material(fresh_state):
    """Unlike the progress bar, the nudge queue should still help the user fill in
    Material name for a process that hasn't been started yet."""
    fresh_state.add_process("Spin Coating", config={"solvents": 0, "solutes": 0})
    dm.rebuild_field_specs(fresh_state)

    queue = dm.build_nudge_queue(fresh_state)

    assert any(item.field_key == "Material name" for item in queue)


def test_build_missing_fields_summary_counts_varying_fields_too(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())

    summary = dm.build_missing_fields_summary(fresh_state)

    assert dict((p.sequence_index, count) for p, count in summary)[process.sequence_index] > 0


def test_build_missing_fields_summary_empty_when_nothing_missing(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    for spec in process.field_specs.values():
        spec.value = "filled"

    assert dm.build_missing_fields_summary(fresh_state) == []


def test_build_nudge_queue_excludes_not_required_missing_fields(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    dm.set_field_required_for_progress(process.field_specs["Recipe file"], False)

    queue = dm.build_nudge_queue(fresh_state)

    assert not any(item.field_key == "Recipe file" for item in queue)


def test_build_nudge_queue_excludes_not_required_outlier_fields(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with(
        "B1", [SPIN_COATING_STEP_EXTREME_OUTLIER, SPIN_COATING_STEP, SPIN_COATING_STEP_CLOSE]
    )
    dm.autofill_process_from_batch(process, "url", "token", cache, "B1", occurrence=0)
    outlier_field = next(
        key for key, spec in process.field_specs.items() if spec.is_outlier and spec.is_filled()
    )
    dm.set_field_required_for_progress(process.field_specs[outlier_field], False)

    queue = dm.build_nudge_queue(fresh_state)

    assert not any(item.field_key == outlier_field for item in queue)


def test_build_missing_fields_summary_excludes_not_required_fields(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    for spec in process.field_specs.values():
        dm.set_field_required_for_progress(spec, False)

    assert dm.build_missing_fields_summary(fresh_state) == []


# ---------------------------------------------------------------------------
# NudgePopupFlow -- light interaction coverage
# ---------------------------------------------------------------------------


def test_nudge_popup_flow_shows_first_missing_item(fresh_state):
    fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    flow = gc.NudgePopupFlow(fresh_state)
    assert len(flow.queue) > 0
    assert len(flow.body.children) == 1  # one item widget shown at a time


def test_nudge_popup_flow_confirm_writes_value_and_advances(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    flow = gc.NudgePopupFlow(fresh_state)
    first_item = flow.queue[0]

    flow._on_confirm(first_item, "confirmed value")

    spec = process.field_specs[first_item.field_key]
    assert spec.value == "confirmed value"
    assert spec.provenance.source == "manual"
    assert flow.index == 1


def test_nudge_popup_flow_confirm_with_blank_value_does_not_write(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    flow = gc.NudgePopupFlow(fresh_state)
    first_item = flow.queue[0]

    flow._on_confirm(first_item, "   ")

    spec = process.field_specs[first_item.field_key]
    assert spec.value is None
    assert flow.index == 1  # still advances, treated like skip


def test_nudge_popup_flow_skip_advances_without_writing(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    flow = gc.NudgePopupFlow(fresh_state)
    first_item = flow.queue[0]

    flow._on_skip()

    spec = process.field_specs[first_item.field_key]
    assert spec.value is None
    assert flow.index == 1


def test_nudge_popup_flow_confirming_outlier_clears_flag_and_sets_manual(fresh_state):
    process = fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 1})
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with(
        "B1", [SPIN_COATING_STEP_EXTREME_OUTLIER, SPIN_COATING_STEP, SPIN_COATING_STEP_CLOSE]
    )
    dm.autofill_process_from_batch(process, "url", "token", cache, "B1", occurrence=0)
    spec = process.field_specs["Annealing temperature [°C]"]
    assert spec.is_outlier is True

    flow = gc.NudgePopupFlow(fresh_state)
    outlier_item = next(item for item in flow.queue if item.kind == "outlier")
    flow._on_confirm(outlier_item, str(spec.value))

    assert spec.is_outlier is False
    assert spec.provenance.source == "manual"


def test_nudge_popup_flow_reaches_summary_after_queue_exhausted(fresh_state):
    fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    flow = gc.NudgePopupFlow(fresh_state)

    for _ in range(len(flow.queue)):
        flow._on_skip()

    summary_html = flow.body.children[0]
    assert "missing" in summary_html.value.lower()


def test_nudge_popup_flow_summary_says_all_filled_when_nothing_missing(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    for spec in process.field_specs.values():
        spec.value = "filled"

    flow = gc.NudgePopupFlow(fresh_state)

    assert "all fields are filled" in flow.body.children[0].value.lower()


def test_nudge_popup_flow_skips_item_filled_elsewhere_since_construction(fresh_state):
    process = fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    flow = gc.NudgePopupFlow(fresh_state)
    first_item = flow.queue[0]

    # simulate the field being filled through a different widget while the flow is open
    process.field_specs[first_item.field_key].value = "filled elsewhere"
    flow._render()

    assert flow.index >= 1  # first item was skipped as no-longer-relevant


def test_nudge_popup_flow_calls_on_change_callback(fresh_state):
    fresh_state.add_process("Laser Scribing")
    dm.rebuild_field_specs(fresh_state)
    calls = []
    flow = gc.NudgePopupFlow(fresh_state, on_change=lambda: calls.append(1))

    flow._on_skip()

    assert calls == [1]


# ---------------------------------------------------------------------------
# Excel finalization: enumerate_sample_rows / compute_nomad_id / resolve_cell_value /
# append_parent_id_column / generate_full_workbook
# ---------------------------------------------------------------------------


def test_enumerate_sample_rows_mother_then_children_uneven_counts(fresh_state):
    fresh_state.add_sample(variation_group_index=0, sample_number=1, child_count=2)
    fresh_state.add_sample(variation_group_index=1, sample_number=2, child_count=0)

    rows = dm.enumerate_sample_rows(fresh_state)

    assert rows == [(1, None), (1, 1), (1, 2), (2, None)]


def _set_id_fields(state, project="JoDa", date="20260526", batch="12"):
    for key, value in (("Project_Name", project), ("Date", date), ("Batch", batch)):
        state.experiment_info_fields[key] = dm.ProcessFieldSpec(key=key, value=value)


def test_compute_nomad_id_kit_format_with_subbatch(fresh_state):
    """KIT_{Project_Name}_{Date}_{Batch}_{Subbatch}_{Sample}, the scheme of Excel_creator's
    Nomad ID formula. Subbatch is computed from the sample's variation Subbatch
    (variation_group_index), 1-based - never a manually-set experiment_info_fields
    ["Subbatch"] value, see subbatch_for_sample."""
    _set_id_fields(fresh_state)
    fresh_state.add_sample(variation_group_index=1, sample_number=3)

    assert dm.compute_nomad_id(fresh_state, sample_number=3, child_index=None) == (
        "KIT_JoDa_20260526_12_2_3"
    )


def test_compute_nomad_id_prefix_comes_from_config(fresh_state):
    from perotf_utils.config import LAB_ID_PREFIX

    _set_id_fields(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    assert dm.compute_nomad_id(fresh_state, 1, None).startswith(f"{LAB_ID_PREFIX}_")


def test_compute_nomad_id_batch_id_is_everything_but_the_sample(fresh_state):
    """The experiment parser derives the batch id as everything before the last "_"."""
    _set_id_fields(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=7)
    nomad_id = dm.compute_nomad_id(fresh_state, 7, None)
    assert "_".join(nomad_id.split("_")[:-1]) == "KIT_JoDa_20260526_12_1"
    assert nomad_id.rsplit("_", 1)[1] == "7"


def test_compute_nomad_id_has_no_child_suffix(fresh_state):
    _set_id_fields(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=3)
    assert dm.compute_nomad_id(fresh_state, 3, child_index=2) == dm.compute_nomad_id(
        fresh_state, 3, child_index=None
    )


def test_compute_nomad_id_omits_empty_subbatch(fresh_state):
    _set_id_fields(fresh_state)

    assert dm.compute_nomad_id(fresh_state, sample_number=1, child_index=None) == (
        "KIT_JoDa_20260526_12_1"
    )


def test_resolve_cell_value_experiment_info_constant_vs_varying(fresh_state):
    fresh_state.experiment_info_fields["Notes"] = dm.ProcessFieldSpec(key="Notes", value="constant")
    assert (
        dm.resolve_cell_value(fresh_state, 0, "Notes", sample_number=1, child_index=None)
        == "constant"
    )

    varying = dm.ProcessFieldSpec(key="Variation", varies=True, per_sample_values={1: "9DMF"})
    fresh_state.experiment_info_fields["Variation"] = varying
    assert (
        dm.resolve_cell_value(fresh_state, 0, "Variation", sample_number=1, child_index=None)
        == "9DMF"
    )
    assert (
        dm.resolve_cell_value(fresh_state, 0, "Variation", sample_number=2, child_index=None)
        is None
    )


def test_resolve_cell_value_pixel_field_same_on_mother_and_child_row(fresh_state):
    # Number of pixels/Pixel area are ordinary Experiment Info fields now (a real exported
    # file showed both as a single constant value across every row of a batch, not blank
    # on a whole-substrate row - see relevant_field_specs' sibling history) - a child row
    # inherits its mother's value same as any other Experiment Info field.
    fresh_state.experiment_info_fields["Number of pixels"] = dm.ProcessFieldSpec(
        key="Number of pixels", value=6
    )
    assert dm.resolve_cell_value(fresh_state, 0, "Number of pixels", 1, child_index=None) == 6
    assert dm.resolve_cell_value(fresh_state, 0, "Number of pixels", 1, child_index=1) == 6


def test_resolve_cell_value_pixel_field_varies_per_sample(fresh_state):
    spec = dm.ProcessFieldSpec(key="Pixel area", varies=True, per_sample_values={1: 0.18, 2: 0.20})
    fresh_state.experiment_info_fields["Pixel area"] = spec
    assert dm.resolve_cell_value(fresh_state, 0, "Pixel area", 1, child_index=None) == 0.18
    assert dm.resolve_cell_value(fresh_state, 0, "Pixel area", 2, child_index=None) == 0.20


def test_resolve_cell_value_process_field_children_inherit_mother_sample_value(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    dm.set_field_varies(process.field_specs["Material name"], True, [1])
    process.field_specs["Material name"].per_sample_values[1] = "C60"

    mother_value = dm.resolve_cell_value(fresh_state, 1, "Material name", 1, child_index=None)
    child_value = dm.resolve_cell_value(fresh_state, 1, "Material name", 1, child_index=3)

    assert mother_value == "C60"
    assert child_value == "C60"  # child inherits the same sample-level value


def test_resolve_cell_value_process_field_constant_broadcasts_to_every_row(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    process.field_specs["Material name"].value = "PCBM"

    assert dm.resolve_cell_value(fresh_state, 1, "Material name", 5, child_index=None) == "PCBM"
    assert dm.resolve_cell_value(fresh_state, 1, "Material name", 7, child_index=2) == "PCBM"


def test_resolve_cell_value_missing_process_returns_none(fresh_state):
    assert dm.resolve_cell_value(fresh_state, 99, "Material name", 1, child_index=None) is None


def test_append_parent_id_column_adds_at_end_without_disturbing_existing_columns(fresh_state):
    fresh_state.add_process("Evaporation")
    workbook = dm.generate_header_workbook(fresh_state)
    worksheet = workbook.active
    original_column_map = dm.build_column_map(worksheet)
    max_col_before = worksheet.max_column

    parent_id_col = dm.append_parent_id_column(worksheet)

    assert parent_id_col == max_col_before + 1
    assert worksheet.cell(row=2, column=parent_id_col).value == "Parent ID"
    # every pre-existing column untouched
    assert dm.build_column_map(worksheet).items() >= original_column_map.items()


def test_generate_full_workbook_has_three_sheets(fresh_state):
    fresh_state.add_process("Evaporation")
    workbook = dm.generate_full_workbook(fresh_state)
    assert set(workbook.sheetnames) == {"Experiment Data", "Data Entry Guide", "How to Cite"}


def test_generate_full_workbook_omits_parent_id_column_when_no_children(fresh_state):
    """Product feedback: with the per-sample child-row UI removed, Parent ID was always a
    fully blank trailing column for the overwhelming majority of real experiments (no
    sample ever has child_count > 0) - it should not appear at all in that case."""
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1, child_count=0)

    workbook = dm.generate_full_workbook(fresh_state)
    worksheet = workbook["Experiment Data"]

    header_values = [
        worksheet.cell(row=2, column=col).value for col in range(1, worksheet.max_column + 1)
    ]
    assert "Parent ID" not in header_values


def test_generate_full_workbook_writes_one_row_per_mother_and_child(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    process.field_specs["Material name"].value = "C60"
    fresh_state.experiment_info_fields["Project_Name"].value = "JoDa"
    fresh_state.experiment_info_fields["Date"].value = "20260526"
    fresh_state.experiment_info_fields["Batch"].value = "12"
    fresh_state.add_sample(variation_group_index=0, sample_number=1, child_count=2)
    fresh_state.add_sample(variation_group_index=0, sample_number=2, child_count=0)

    workbook = dm.generate_full_workbook(fresh_state)
    worksheet = workbook["Experiment Data"]
    column_map = dm.build_column_map(worksheet)

    nomad_id_col = column_map[(0, "Nomad ID")]
    sample_col = column_map[(0, "Sample")]
    subbatch_col = column_map[(0, "Subbatch")]
    parent_id_col = worksheet.max_column  # appended last

    # both samples are in variation Subbatch 0 (1-based Subbatch value "1")
    assert worksheet.cell(row=3, column=nomad_id_col).value == "KIT_JoDa_20260526_12_1_1"
    assert worksheet.cell(row=3, column=sample_col).value == 1
    assert worksheet.cell(row=3, column=subbatch_col).value == "1"
    assert worksheet.cell(row=3, column=parent_id_col).value is None  # mother

    # peroTF ids have no child segment: a child row carries its mother's id
    assert worksheet.cell(row=4, column=nomad_id_col).value == "KIT_JoDa_20260526_12_1_1"
    assert worksheet.cell(row=4, column=parent_id_col).value == "KIT_JoDa_20260526_12_1_1"

    assert worksheet.cell(row=6, column=nomad_id_col).value == "KIT_JoDa_20260526_12_1_2"
    assert worksheet.cell(row=6, column=sample_col).value == 2
    assert worksheet.cell(row=6, column=subbatch_col).value == "1"

    material_col = column_map[(process.sequence_index, "Material name")]
    assert worksheet.cell(row=3, column=material_col).value == "C60"
    assert worksheet.cell(row=4, column=material_col).value == "C60"  # child inherits


def test_generate_full_workbook_pixel_fields_written_on_mother_and_inherited_by_child(
    fresh_state,
):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.experiment_info_fields["Number of pixels"].value = 6
    fresh_state.add_sample(variation_group_index=0, sample_number=1, child_count=1)

    workbook = dm.generate_full_workbook(fresh_state)
    worksheet = workbook["Experiment Data"]
    column_map = dm.build_column_map(worksheet)
    pixels_col = column_map[(0, "Number of pixels")]

    assert worksheet.cell(row=3, column=pixels_col).value == 6  # mother
    assert worksheet.cell(row=4, column=pixels_col).value == 6  # child inherits


def test_workbook_to_bytes_returns_valid_xlsx_signature(fresh_state):
    fresh_state.add_process("Evaporation")
    workbook = dm.generate_full_workbook(fresh_state)
    data = dm.workbook_to_bytes(workbook)
    assert data.startswith(b"PK")
    assert len(data) > 0


def test_build_experiment_filename_matches_date_convention():
    from datetime import datetime

    filename = dm.build_experiment_filename()
    assert filename == f"{datetime.now().strftime('%Y%m%d')}_experiment_file.xlsx"


# ---------------------------------------------------------------------------
# ExperimentInfoPanel
# ---------------------------------------------------------------------------


def test_experiment_info_panel_excludes_only_computed_keys(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    panel = gc.ExperimentInfoPanel(fresh_state)

    # panel.children[0] is the provenance summary, not a field row - skip it
    rendered_labels = {row.children[1].value for row in panel.children[1:]}

    assert "Variation" not in rendered_labels
    assert "Nomad ID" not in rendered_labels
    assert "Sample" not in rendered_labels
    assert "Subbatch" not in rendered_labels
    assert "Project_Name *" in rendered_labels  # a normal (required) field is still shown
    # Number of pixels/Pixel area are ordinary Experiment Info fields now, not excluded -
    # see relevant_field_specs' sibling history for why the old per-CHILD-row gating was
    # removed.
    assert "Number of pixels *" in rendered_labels
    assert "Pixel area *" in rendered_labels


def test_experiment_info_panel_value_edit_sets_manual_provenance(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    panel = gc.ExperimentInfoPanel(fresh_state)

    panel._on_value_change("Project_Name", "CsFA")

    spec = fresh_state.experiment_info_fields["Project_Name"]
    assert spec.value == "CsFA"
    assert spec.provenance.source == "manual"


def test_experiment_info_panel_varies_checkbox_promotes_scope(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    fresh_state.experiment_info_fields["Sample dimension"].value = "1 cm x 1 cm"
    panel = gc.ExperimentInfoPanel(fresh_state)

    panel._on_varies_change("Sample dimension", True)

    spec = fresh_state.experiment_info_fields["Sample dimension"]
    assert spec.varies is True
    assert spec.per_sample_values == {1: "1 cm x 1 cm"}


def test_experiment_info_panel_calls_on_change_callback(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    calls = []
    panel = gc.ExperimentInfoPanel(fresh_state, on_change=lambda: calls.append(1))

    panel._on_value_change("Project_Name", "CsFA")

    assert calls == [1]


def test_experiment_info_panel_not_required_field_has_no_asterisk(fresh_state):
    """The per-row 'Required' checkbox was removed (required_for_progress is now
    config/required_fields.json-driven) - "Notes" is a global exception there, so its
    label never gets a '*'."""
    dm.rebuild_field_specs(fresh_state)
    panel = gc.ExperimentInfoPanel(fresh_state)

    rendered_labels = {row.children[1].value for row in panel.children[1:]}

    assert "Notes" in rendered_labels
    assert "Notes *" not in rendered_labels


# ---------------------------------------------------------------------------
# SampleSetupPanel
# ---------------------------------------------------------------------------


def test_sample_setup_panel_defaults_are_16_samples_4_subbatches(fresh_state):
    panel = gc.SampleSetupPanel(fresh_state)
    assert panel.total_samples_input.value == 16
    assert panel.set_count_input.value == 4


def test_sample_setup_panel_apply_adds_samples_for_single_set(fresh_state):
    panel = gc.SampleSetupPanel(fresh_state)
    panel.set_count_input.value = 1
    panel.sets_inputs_box.children[0].value = 3

    panel._on_apply(None)

    assert len(fresh_state.samples) == 3
    assert all(s.variation_group_index == 0 for s in fresh_state.samples)
    assert [s.sample_number for s in fresh_state.samples] == [1, 2, 3]


def test_sample_setup_panel_apply_sample_setup_is_public_equivalent_of_apply_click(
    fresh_state,
):
    panel = gc.SampleSetupPanel(fresh_state)
    panel.set_count_input.value = 1
    panel.sets_inputs_box.children[0].value = 3

    panel.apply_sample_setup()  # same action app.py calls once on page load

    assert len(fresh_state.samples) == 3


def test_sample_setup_panel_apply_removes_excess_samples_when_count_lowered(fresh_state):
    panel = gc.SampleSetupPanel(fresh_state)
    panel.set_count_input.value = 1
    panel.sets_inputs_box.children[0].value = 3
    panel._on_apply(None)
    fresh_state.samples[0].child_count = 2  # simulate user-configured data on sample 1

    # user lowers the requested count and re-applies
    panel.sets_inputs_box.children[0].value = 1
    panel._on_apply(None)

    assert len(fresh_state.samples) == 1  # excess samples removed, not just left to accumulate
    assert fresh_state.samples[0].sample_number == 1  # earliest sample kept, not the latest
    assert fresh_state.samples[0].child_count == 2  # its data survives the shrink


def test_sample_setup_panel_lowering_total_samples_shrinks_after_apply(fresh_state):
    """Regression test: changing Total samples used to clamp each Subbatch's shown value
    to max(existing_count, new_default), so a lower total was silently ignored and
    re-applying could only ever grow the batch, never shrink it back down."""
    panel = gc.SampleSetupPanel(fresh_state)
    panel.set_count_input.value = 1
    panel.sets_inputs_box.children[0].value = 5
    panel._on_apply(None)
    assert len(fresh_state.samples) == 5

    panel.total_samples_input.value = 2

    assert panel.sets_inputs_box.children[0].value == 2  # no longer clamped up to the old 5

    panel._on_apply(None)

    assert len(fresh_state.samples) == 2


def test_sample_setup_panel_lowering_subbatch_count_removes_orphaned_samples(fresh_state):
    """Regression test: lowering Variation Subbatch count used to leave samples from the
    now-unconfigured Subbatches stranded in state forever, with no row left to remove them."""
    panel = gc.SampleSetupPanel(fresh_state)  # defaults: 16 samples / 4 subbatches
    panel._on_apply(None)
    assert {s.variation_group_index for s in fresh_state.samples} == {0, 1, 2, 3}

    panel.set_count_input.value = 2

    panel._on_apply(None)

    assert {s.variation_group_index for s in fresh_state.samples} == {0, 1}


def test_sample_setup_panel_multiple_sets_uneven_counts(fresh_state):
    panel = gc.SampleSetupPanel(fresh_state)
    panel.set_count_input.value = 2
    panel.sets_inputs_box.children[0].value = 3
    panel.sets_inputs_box.children[1].value = 5

    panel._on_apply(None)

    set_0 = [s for s in fresh_state.samples if s.variation_group_index == 0]
    set_1 = [s for s in fresh_state.samples if s.variation_group_index == 1]
    assert len(set_0) == 3
    assert len(set_1) == 5


def test_sample_setup_panel_preloads_natural_division(fresh_state):
    panel = gc.SampleSetupPanel(fresh_state)
    panel.set_count_input.value = 4
    panel.total_samples_input.value = 15

    assert [c.value for c in panel.sets_inputs_box.children] == [4, 4, 4, 3]


def test_sample_setup_panel_has_no_samples_table(fresh_state):
    """The per-sample list/remove-button table was removed per product feedback
    ("doesn't make sense here") - ExperimentState.remove_sample still exists for
    programmatic use, just not exposed from this panel."""
    panel = gc.SampleSetupPanel(fresh_state)
    assert not hasattr(panel, "samples_table")
    assert not hasattr(panel, "_on_remove_sample")


def test_sample_setup_panel_calls_on_change_callback(fresh_state):
    calls = []
    panel = gc.SampleSetupPanel(fresh_state, on_change=lambda: calls.append(1))

    panel._on_apply(None)

    assert calls == [1]


# ---------------------------------------------------------------------------
# create_download_button
# ---------------------------------------------------------------------------


def test_create_download_button_click_produces_download_link(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    section = gc.create_download_button(fresh_state)
    button, output_area = section.children

    button.click()

    assert "download" in output_area.value.lower()
    assert "base64," in output_area.value
    assert ".xlsx" in output_area.value


# ---------------------------------------------------------------------------
# create_quick_fill_all_button
# ---------------------------------------------------------------------------


def test_create_quick_fill_all_button_fills_and_calls_on_change(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    _add_manual_field(process, "Operator")
    calls = []

    section = gc.create_quick_fill_all_button(fresh_state, on_change=lambda: calls.append(1))
    button, status = section.children
    button.click()

    assert fresh_state.experiment_info_fields["Date"].value
    assert process.field_specs["Operator"].value
    assert "date field" in status.value.lower()
    assert calls == [1]


# ---------------------------------------------------------------------------
# upload_experiment_excel -- mocked request-shape tests only. NOT verified against the
# real API by this suite - see tests/live/test_smart_databaser_upload.py, which must be
# run manually against a disposable upload before this is trusted in production.
# ---------------------------------------------------------------------------


def _mock_response(json_data=None, raise_error=False):
    response = MagicMock()
    if raise_error:
        response.raise_for_status.side_effect = requests.HTTPError("boom")
    else:
        response.raise_for_status.side_effect = None
    if json_data is not None:
        response.json.return_value = json_data
    return response


def test_upload_experiment_excel_put_uses_correct_url_headers_and_file_tuple():
    with (
        patch.object(dm.requests, "put", return_value=_mock_response()) as mock_put,
        patch.object(dm.requests, "post", return_value=_mock_response()),
        patch.object(
            dm.requests,
            "get",
            return_value=_mock_response({"data": {"process_running": False}}),
        ),
        patch.object(dm.time, "sleep"),
    ):
        dm.upload_experiment_excel(
            "https://nomad.example/api/v1", "tok", "UP1", "test.xlsx", b"bytes"
        )

    mock_put.assert_called_once()
    args, kwargs = mock_put.call_args
    assert args[0] == "https://nomad.example/api/v1/uploads/UP1/raw/"
    assert kwargs["headers"] == {"Authorization": "Bearer tok"}
    assert kwargs["data"] == {"wait_for_processing": False}
    filename, file_bytes, mime = kwargs["files"]["file"]
    assert filename == "test.xlsx"
    assert file_bytes == b"bytes"
    assert mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_upload_experiment_excel_triggers_process_action_after_put():
    with (
        patch.object(dm.requests, "put", return_value=_mock_response()),
        patch.object(dm.requests, "post", return_value=_mock_response()) as mock_post,
        patch.object(
            dm.requests,
            "get",
            return_value=_mock_response({"data": {"process_running": False}}),
        ),
        patch.object(dm.time, "sleep"),
    ):
        dm.upload_experiment_excel(
            "https://nomad.example/api/v1", "tok", "UP1", "test.xlsx", b"bytes"
        )

    mock_post.assert_called_once_with(
        "https://nomad.example/api/v1/uploads/UP1/action/process",
        headers={"Authorization": "Bearer tok"},
    )


def test_upload_experiment_excel_polls_until_processing_finishes():
    poll_responses = [
        _mock_response({"data": {"process_running": True}}),
        _mock_response({"data": {"process_running": True}}),
        _mock_response({"data": {"process_running": False}}),
    ]
    with (
        patch.object(dm.requests, "put", return_value=_mock_response()),
        patch.object(dm.requests, "post", return_value=_mock_response()),
        patch.object(dm.requests, "get", side_effect=poll_responses) as mock_get,
        patch.object(dm.time, "sleep"),
    ):
        dm.upload_experiment_excel(
            "https://nomad.example/api/v1", "tok", "UP1", "test.xlsx", b"bytes"
        )

    assert mock_get.call_count == 3
    mock_get.assert_called_with(
        "https://nomad.example/api/v1/uploads/UP1", headers={"Authorization": "Bearer tok"}
    )


def test_upload_experiment_excel_raises_on_put_failure():
    with (
        patch.object(dm.requests, "put", return_value=_mock_response(raise_error=True)),
        patch.object(dm.requests, "post", return_value=_mock_response()),
        patch.object(dm.requests, "get"),
        patch.object(dm.time, "sleep"),
        pytest.raises(requests.HTTPError),
    ):
        dm.upload_experiment_excel(
            "https://nomad.example/api/v1", "tok", "UP1", "test.xlsx", b"bytes"
        )


def test_upload_experiment_excel_ignores_process_action_failure_and_still_polls():
    """Seen live: POST .../action/process can return 400 even though the upload
    proceeds and finishes normally - likely because PUTting the file already triggers
    processing automatically on this NOMAD version. Must not raise here - the polling
    loop is the real source of truth."""
    with (
        patch.object(dm.requests, "put", return_value=_mock_response()),
        patch.object(dm.requests, "post", return_value=_mock_response(raise_error=True)),
        patch.object(
            dm.requests,
            "get",
            return_value=_mock_response({"data": {"process_running": False}}),
        ) as mock_get,
        patch.object(dm.time, "sleep"),
    ):
        dm.upload_experiment_excel(
            "https://nomad.example/api/v1", "tok", "UP1", "test.xlsx", b"bytes"
        )

    mock_get.assert_called()


def test_upload_experiment_excel_raises_timeout_if_never_finishes():
    with (
        patch.object(dm.requests, "put", return_value=_mock_response()),
        patch.object(dm.requests, "post", return_value=_mock_response()),
        patch.object(
            dm.requests,
            "get",
            return_value=_mock_response({"data": {"process_running": True}}),
        ),
        patch.object(dm.time, "sleep"),
        pytest.raises(TimeoutError),
    ):
        dm.upload_experiment_excel(
            "https://nomad.example/api/v1",
            "tok",
            "UP1",
            "test.xlsx",
            b"bytes",
            poll_interval_seconds=1.0,
            max_poll_seconds=2.0,
        )


# ---------------------------------------------------------------------------
# create_finish_section -- three explicit end-of-workflow actions
# ---------------------------------------------------------------------------


def _fill_critical_fields(state):
    """Batch/Project_Name/Date are hard-required for every create_finish_section action (see
    missing_critical_fields) - tests exercising OTHER mechanics (nudge gate, upload
    call, ...) need these filled first or every action is blocked before it even starts.
    See test_create_finish_section_blocks_* below for tests of the block itself."""
    state.experiment_info_fields["Batch"].value = "1"
    state.experiment_info_fields["Project_Name"].value = "Test"
    state.experiment_info_fields["Date"].value = "20260526"


def test_create_finish_section_download_only_produces_link(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    _fill_critical_fields(fresh_state)
    cache = dm.NomadSessionCache()

    with patch.object(
        dm, "get_all_uploads", return_value=[{"upload_id": "UP1", "upload_name": "Test"}]
    ):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            upload_dropdown,
            buttons_row,
            _nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        skip_checkbox.value = True  # bypass the nudge gate for this mechanics-only test
        download_button, _upload_button, _combo_button = buttons_row.children
        download_button.click()

    assert "download" in status_output.value.lower()
    assert "base64," in status_output.value


def test_create_finish_section_upload_only_without_target_shows_error(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    _fill_critical_fields(fresh_state)
    cache = dm.NomadSessionCache()

    with patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            _upload_dropdown,
            buttons_row,
            _nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        skip_checkbox.value = True
        _download_button, upload_button, _combo_button = buttons_row.children
        upload_button.click()

    assert "target upload" in status_output.value.lower()


def test_create_finish_section_upload_only_calls_upload_experiment_excel(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    _fill_critical_fields(fresh_state)
    cache = dm.NomadSessionCache()

    with (
        patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]),
        patch.object(gc, "upload_experiment_excel") as mock_upload,
    ):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            upload_dropdown,
            buttons_row,
            _nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        skip_checkbox.value = True
        upload_dropdown.value = "UP1"
        _download_button, upload_button, _combo_button = buttons_row.children
        upload_button.click()

    mock_upload.assert_called_once()
    assert mock_upload.call_args[0][2] == "UP1"
    assert "Uploaded" in status_output.value


def test_create_finish_section_download_and_upload_does_both(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    _fill_critical_fields(fresh_state)
    cache = dm.NomadSessionCache()

    with (
        patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]),
        patch.object(gc, "upload_experiment_excel") as mock_upload,
    ):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            upload_dropdown,
            buttons_row,
            _nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        skip_checkbox.value = True
        upload_dropdown.value = "UP1"
        _download_button, _upload_button, combo_button = buttons_row.children
        combo_button.click()

    mock_upload.assert_called_once()
    assert "base64," in status_output.value
    assert "Uploaded" in status_output.value


def test_create_finish_section_download_and_upload_reports_upload_failure(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    _fill_critical_fields(fresh_state)
    cache = dm.NomadSessionCache()

    with (
        patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]),
        patch.object(gc, "upload_experiment_excel", side_effect=RuntimeError("network down")),
    ):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            upload_dropdown,
            buttons_row,
            _nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        skip_checkbox.value = True
        upload_dropdown.value = "UP1"
        _download_button, _upload_button, combo_button = buttons_row.children
        combo_button.click()

    assert "base64," in status_output.value  # download still succeeded
    assert "failed" in status_output.value.lower()


def test_create_finish_section_gates_download_behind_nudge_flow_by_default(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    _fill_critical_fields(fresh_state)
    cache = dm.NomadSessionCache()

    with patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            _upload_dropdown,
            buttons_row,
            nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        assert skip_checkbox.value is False
        download_button, _upload_button, _combo_button = buttons_row.children
        download_button.click()

        # nudge flow opened; the download hasn't happened yet
        assert len(nudge_area.children) == 2
        assert status_output.value == ""

        continue_button = nudge_area.children[1]
        continue_button.click()

    assert "base64," in status_output.value
    assert nudge_area.children == ()


def test_create_finish_section_skip_checkbox_bypasses_nudge_flow(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    _fill_critical_fields(fresh_state)
    cache = dm.NomadSessionCache()

    with patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            _upload_dropdown,
            buttons_row,
            nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        skip_checkbox.value = True
        download_button, _upload_button, _combo_button = buttons_row.children
        download_button.click()

    assert nudge_area.children == ()
    assert "base64," in status_output.value


def test_create_finish_section_blocks_download_when_batch_missing(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.experiment_info_fields["Project_Name"].value = "Test"  # Batch left empty
    cache = dm.NomadSessionCache()

    with patch.object(dm, "get_all_uploads", return_value=[]):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            _upload_dropdown,
            buttons_row,
            nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        skip_checkbox.value = True  # must block even with the nudge gate skipped
        download_button, _upload_button, _combo_button = buttons_row.children
        download_button.click()

    assert "Batch" in status_output.value
    assert "base64," not in status_output.value
    assert nudge_area.children == ()


def test_create_finish_section_blocks_upload_when_project_name_missing(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.experiment_info_fields["Batch"].value = "1"  # Project_Name left empty
    cache = dm.NomadSessionCache()

    with (
        patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]),
        patch.object(gc, "upload_experiment_excel") as mock_upload,
    ):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            upload_dropdown,
            buttons_row,
            _nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        skip_checkbox.value = True
        upload_dropdown.value = "UP1"
        _download_button, upload_button, _combo_button = buttons_row.children
        upload_button.click()

    assert "Project_Name" in status_output.value
    mock_upload.assert_not_called()


def test_create_finish_section_blocks_both_missing_lists_both_fields(fresh_state):
    fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    cache = dm.NomadSessionCache()

    with patch.object(dm, "get_all_uploads", return_value=[]):
        section = gc.create_finish_section(fresh_state, "url", "token", cache)
        (
            skip_checkbox,
            _caption,
            _upload_dropdown,
            buttons_row,
            _nudge_area,
            status_output,
            _download_js_output,
        ) = section.children
        skip_checkbox.value = True
        download_button, _upload_button, _combo_button = buttons_row.children
        download_button.click()

    assert "Batch" in status_output.value
    assert "Project_Name" in status_output.value


# ---------------------------------------------------------------------------
# subbatch_for_sample / required_fields.json / Experiment Info autofill / field-value
# multipliers / field-mapping debug report - all new 2026-07-28.
# ---------------------------------------------------------------------------


def test_subbatch_for_sample_is_one_based_variation_group(fresh_state):
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    fresh_state.add_sample(variation_group_index=2, sample_number=2)

    assert dm.subbatch_for_sample(fresh_state, 1) == "1"
    assert dm.subbatch_for_sample(fresh_state, 2) == "3"


def test_subbatch_for_sample_missing_sample_returns_none(fresh_state):
    assert dm.subbatch_for_sample(fresh_state, 99) is None


def test_is_field_required_by_default_global_exception():
    assert dm.is_field_required_by_default("Laser Scribing", "Notes") is False
    assert dm.is_field_required_by_default("Experiment Info", "Notes") is False


def test_is_field_required_by_default_true_for_ordinary_field():
    assert dm.is_field_required_by_default("Laser Scribing", "Recipe file") is True


def test_load_required_field_exceptions_matches_shipped_config():
    global_exceptions, by_process_type = dm.load_required_field_exceptions()
    assert "Notes" in global_exceptions
    assert isinstance(by_process_type, dict)


def test_sync_field_specs_from_columns_uses_required_fields_config(fresh_state):
    process = fresh_state.add_process("Annealing")
    dm.rebuild_field_specs(fresh_state)
    assert process.field_specs["Notes"].required_for_progress is False
    assert process.field_specs["Annealing time [min]"].required_for_progress is True


def test_load_field_value_multipliers_confirms_cleaning_time_conversion():
    multipliers = build_field_value_multipliers()
    assert multipliers["Cleaning UV-Ozone"]["UV-Ozone Time [s]"] == 60
    assert multipliers["Cleaning UV-Ozone"]["Time 1 [s]"] == 60
    assert multipliers["Cleaning O2-Plasma"]["Gas-Plasma Time [s]"] == 60


def test_fetch_experiment_info_source_follows_substrate_reference():
    cache = dm.NomadSessionCache()
    sample_data = {"number_of_junctions": 1, "substrate": "../uploads/U1/archive/SUB1#/data"}
    substrate_data = {"substrate": "Soda Lime Glass", "conducting_material": ["ITO"]}

    with (
        patch.object(dm, "get_ids_in_batch", return_value=["S1"]),
        patch.object(dm, "get_entryid", return_value="ENTRY1"),
        patch.object(
            dm,
            "get_entry_data",
            side_effect=lambda url, token, entry_id: (
                sample_data if entry_id == "ENTRY1" else substrate_data
            ),
        ),
    ):
        source = dm.fetch_experiment_info_source("url", "token", cache, "B1")

    assert source == {"sample": sample_data, "substrate": substrate_data}


def test_fetch_experiment_info_source_no_samples_returns_none():
    cache = dm.NomadSessionCache()
    with patch.object(dm, "get_ids_in_batch", return_value=[]):
        assert dm.fetch_experiment_info_source("url", "token", cache, "B1") is None


def test_autofill_experiment_info_from_batch_fills_mapped_fields_skips_never_autofilled(
    fresh_state,
):
    dm.rebuild_field_specs(fresh_state)
    source = {
        "sample": {"number_of_junctions": 2},
        "substrate": {
            "substrate": "Soda Lime Glass",
            "conducting_material": ["ITO"],
            "solar_cell_area": 6.4516,
        },
    }
    cache = _cache_with("B1", [], experiment_info_source=source)

    written = dm.autofill_experiment_info_from_batch(fresh_state, "url", "token", cache, "B1")

    assert written > 0
    assert fresh_state.experiment_info_fields["Number of junctions"].value == 2
    assert fresh_state.experiment_info_fields["Substrate material"].value == "Soda Lime Glass"
    assert fresh_state.experiment_info_fields["Substrate conductive layer"].value == "ITO"
    assert fresh_state.experiment_info_fields["Sample area [cm^2]"].value == 6.4516
    # never autofilled, even though nothing in `source` maps to them anyway
    assert fresh_state.experiment_info_fields["Date"].value is None
    assert fresh_state.experiment_info_fields["Project_Name"].value is None
    assert fresh_state.experiment_info_fields["Batch"].value is None


def test_autofill_experiment_info_from_batch_no_source_writes_nothing(fresh_state):
    dm.rebuild_field_specs(fresh_state)
    cache = _cache_with("B1", [], experiment_info_source=None)

    written = dm.autofill_experiment_info_from_batch(fresh_state, "url", "token", cache, "B1")

    assert written == 0


def test_apply_whole_experiment_template_also_autofills_experiment_info(fresh_state):
    source = {"sample": {"number_of_junctions": 1}, "substrate": {"substrate": "Soda Lime Glass"}}
    cache = _cache_with("B1", [SPIN_COATING_STEP], experiment_info_source=source)

    dm.apply_whole_experiment_template(fresh_state, "url", "token", cache, "B1")

    assert fresh_state.experiment_info_fields["Substrate material"].value == "Soda Lime Glass"
    assert fresh_state.experiment_info_fields["Project_Name"].value is None


def test_build_field_mapping_debug_report_reports_mapped_and_ignored():
    report = dm.build_field_mapping_debug_report("Cleaning UV-Ozone", CLEANING_STEP)

    mapped_by_key = {row["excel_key"]: row for row in report["mapped"]}
    assert mapped_by_key["UV-Ozone Time [s]"]["value"] == 900.0
    assert mapped_by_key["Solvent 1"]["value"] == "Hellmanex-DI water"

    ignored_paths = {row["path"] for row in report["ignored"]}
    # "method"/"positon_in_experimental_plan" are real raw fields no mapping claims
    assert "method" in ignored_paths
    assert "positon_in_experimental_plan" in ignored_paths


def test_build_field_mapping_debug_report_unmapped_process_type_reports_everything_ignored():
    """Multijunction Info has no process_specs.py archive path at all (the parser
    creates no entry for it) and no _DERIVED_FIELDS entry either."""
    report = dm.build_field_mapping_debug_report("Multijunction Info", {"method": "Cleaning"})
    assert report["mapped"] == []
    assert {row["path"] for row in report["ignored"]} == {"method"}


def test_batch_field_mapping_debug_panel_loads_batch_and_renders_report(fresh_state):
    cache = dm.NomadSessionCache()
    with (
        patch.object(dm, "get_batch_ids", return_value=["B1"]),
        patch.object(dm, "get_ids_in_batch", return_value=["S1"]),
        patch.object(dm, "get_processing_steps", return_value=[SPIN_COATING_STEP]),
    ):
        panel = gc.BatchFieldMappingDebugPanel("url", "token", cache)
        _caption, batch_picker, process_type_dropdown, _occurrence_dropdown, report_output = (
            panel.children
        )
        process_type_dropdown.value = "Spin Coating"
        search_field, selector, load_button, status = batch_picker.children
        selector.value = ("B1",)
        load_button.click()

    assert "Loaded B1" in status.value
    assert "Material name" in report_output.value
    assert "Me4PACz" in report_output.value


def test_build_field_row_date_field_today_button_fills_current_date(fresh_state):
    from datetime import datetime

    fresh_state.experiment_info_fields["Date"] = dm.ProcessFieldSpec(key="Date")
    panel = gc.ExperimentInfoPanel(fresh_state)

    rows_by_label = {row.children[1].value: row for row in panel.children[1:]}
    date_row = rows_by_label["Date *"]
    value_widget, today_button = date_row.children[2], date_row.children[3]
    assert today_button.description == "Today"

    today_button.click()

    # YYYYMMDD: the Date is a segment of every Nomad ID
    assert value_widget.value == datetime.now().strftime("%Y%m%d")


def _standalone_field_row(field_key):
    """A _FieldRow for a hand-made spec, recording value changes. peroTF's Excel_creator
    writes no Datetime/Operator column, so no process panel renders one; the quick-fill
    and guard logic of the row itself is still tested here."""
    spec = dm.ProcessFieldSpec(key=field_key)
    changes = []
    row = gc._FieldRow(
        field_key,
        spec,
        on_varies_change=lambda key, varies: None,
        on_value_change=lambda key, value: changes.append((key, value)),
    )
    return row, changes


def test_build_field_row_datetime_field_colon_is_not_a_forbidden_character():
    """Regression test: Datetime's "%d.%m.%Y %H:%M:%S" format needs colons, and this
    field never feeds compute_nomad_id/build_experiment_filename - it must be exempt from
    the forbidden-character guard (colon is otherwise in FORBIDDEN_VALUE_CHARACTERS)."""
    datetime_row, changes = _standalone_field_row("Datetime")
    value_widget, today_button, warning = (
        datetime_row.children[2],
        datetime_row.children[3],
        datetime_row.children[-1],
    )
    assert today_button.description == "Today"

    today_button.click()

    assert ":" in value_widget.value
    assert changes == [("Datetime", value_widget.value)]
    assert warning.value == ""
    assert not value_widget.layout.border  # guard never even touched this widget


def test_build_field_row_operator_field_me_button_fills_current_user(monkeypatch):
    monkeypatch.setenv("NOMAD_CLIENT_USER", "Jane Doe")
    operator_row, _changes = _standalone_field_row("Operator")
    value_widget, me_button = operator_row.children[2], operator_row.children[3]
    assert me_button.description == "Me"

    me_button.click()

    assert value_widget.value == "Jane Doe"


# ---------------------------------------------------------------------------
# Custom Variation template ("\1"/"\2"/... syntax) and the forbidden-character guard on
# every data-value Text widget - both added 2026-07-28.
# ---------------------------------------------------------------------------


def test_render_variation_template_substitutes_by_position():
    assert dm.render_variation_template(r"Den=\1_Sol-\2", ["1", "ipa"]) == "Den=1_Sol-ipa"


def test_render_variation_template_not_every_value_has_to_be_used():
    # \2 is skipped entirely - the template just never references it
    assert dm.render_variation_template(r"Sol-\2", ["1", "ipa", "20"]) == "Sol-ipa"


def test_render_variation_template_missing_value_becomes_blank():
    assert dm.render_variation_template(r"Den=\1_Sol-\2", ["1", None]) == "Den=1_Sol-"


def test_render_variation_template_out_of_range_index_becomes_blank():
    assert dm.render_variation_template(r"Den=\1_Extra-\5", ["1"]) == "Den=1_Extra-"


def test_compute_variation_label_uses_template_when_set(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 1, "solutes": 0})
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.get_process(1)
    dm.set_field_varies(process.field_specs["Material name"], True, [])
    process.field_specs["Material name"].per_sample_values[1] = "PCBM"
    fresh_state.variation_template = r"Mat=\1"

    assert dm.compute_variation_label(fresh_state, sample_number=1) == "Mat=PCBM"


def test_apply_variation_template_regenerates_computed_values_only(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 0, "solutes": 0})
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.get_process(1)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    fresh_state.add_sample(variation_group_index=0, sample_number=2)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())
    process.field_specs["Material name"].per_sample_values = {1: "PCBM", 2: "Spiro"}
    dm.update_variation_column(fresh_state)  # seeds the automatic label first
    # sample 2's Variation was manually overridden - must survive a template switch
    fresh_state.experiment_info_fields["Variation"].per_sample_provenance[2] = dm.FieldProvenance(
        source="manual"
    )

    dm.apply_variation_template(fresh_state, r"Mat=\1")

    assert fresh_state.variation_template == r"Mat=\1"
    variation_spec = fresh_state.experiment_info_fields["Variation"]
    assert variation_spec.per_sample_values[1] == "Mat=PCBM"
    # "manual"-provenance value untouched by the template switch
    assert variation_spec.per_sample_values[2] == "material-name-Spiro"


def test_apply_variation_template_blank_reverts_to_automatic_label(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 0, "solutes": 0})
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.get_process(1)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())
    process.field_specs["Material name"].per_sample_values[1] = "PCBM"
    dm.apply_variation_template(fresh_state, r"Mat=\1")

    dm.apply_variation_template(fresh_state, "   ")

    assert fresh_state.variation_template is None
    variation_spec = fresh_state.experiment_info_fields["Variation"]
    assert variation_spec.per_sample_values[1] == "material-name-PCBM"


def test_variation_template_panel_apply_updates_state_and_legend(fresh_state):
    fresh_state.add_process("Spin Coating", config={"solvents": 0, "solutes": 0})
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.get_process(1)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    dm.set_field_varies(process.field_specs["Material name"], True, fresh_state.sample_numbers())
    process.field_specs["Material name"].per_sample_values[1] = "PCBM"
    calls = []
    panel = gc.VariationTemplatePanel(fresh_state, on_change=lambda: calls.append(1))
    assert r"\1" in panel.legend.value  # legend shows the current \N mapping

    panel.template_input.value = r"Mat=\1"
    panel.apply_button.click()

    assert fresh_state.variation_template == r"Mat=\1"
    assert fresh_state.experiment_info_fields["Variation"].per_sample_values[1] == "Mat=PCBM"
    assert calls == [1]


def test_variation_template_panel_buttons_are_labeled_positioned_and_colored(fresh_state):
    """Product ask: 'Apply Formula' (not just 'Apply', to make clear it's formula-based)
    sits right next to 'Automatically fill up Variation' (the pre-established/default
    fill), and both are color-coded so they're easy to tell apart at a glance."""
    panel = gc.VariationTemplatePanel(fresh_state)

    assert panel.apply_button.description == "Apply Formula"
    assert panel.apply_button.button_style == "primary"
    assert panel.autofill_button.description == "Automatically fill up Variation"
    assert panel.autofill_button.button_style == "success"

    button_row = panel.children[1]
    assert list(button_row.children) == [
        panel.template_input,
        panel.apply_button,
        panel.autofill_button,
    ]


def test_variation_template_panel_hides_when_no_matrix_table(fresh_state):
    """Same show/hide condition as VaryingFieldsMatrix's own placeholder - no point
    offering a custom Variation format when there's no varying-fields table to apply it
    to (no varying field, or no sample yet)."""
    fresh_state.add_process("Spin Coating", config={"solvents": 0, "solutes": 0})
    dm.rebuild_field_specs(fresh_state)
    process = fresh_state.get_process(1)
    panel = gc.VariationTemplatePanel(fresh_state)
    assert panel.layout.display == "none"

    # a varying field alone, with no sample, still isn't a real table
    dm.set_field_varies(process.field_specs["Material name"], True, [])
    panel.refresh()
    assert panel.layout.display == "none"

    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    panel.refresh()

    assert panel.layout.display == ""
    assert "Material name" in panel.legend.value


def test_find_forbidden_characters_detects_reserved_set():
    assert dm.find_forbidden_characters(r"bad\name") == ["\\"]
    assert (
        set(dm.find_forbidden_characters('a/b:c*d?e"f<g>h|i\\j')) == dm.FORBIDDEN_VALUE_CHARACTERS
    )
    assert dm.find_forbidden_characters("normal-value_1.5") == []
    assert dm.find_forbidden_characters(None) == []


def test_build_field_row_rejects_forbidden_character_value(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    panel = gc.ProcessFieldsPanel(fresh_state, process)

    rows_by_label = {row.children[1].value: row for row in panel.children[1:]}
    row = rows_by_label["Material name *"]
    value_widget, warning = row.children[2], row.children[-1]

    value_widget.value = "bad\\name"

    assert process.field_specs["Material name"].value is None  # never persisted
    assert "\\" in warning.value
    assert value_widget.layout.border != ""

    value_widget.value = "good-name"

    assert process.field_specs["Material name"].value == "good-name"
    assert warning.value == ""


def test_varying_fields_matrix_rejects_forbidden_character_cell(fresh_state):
    process = fresh_state.add_process("Evaporation")
    dm.rebuild_field_specs(fresh_state)
    fresh_state.add_sample(variation_group_index=0, sample_number=1)
    spec = process.field_specs["Material name"]
    dm.set_field_varies(spec, True, fresh_state.sample_numbers())

    matrix = gc.VaryingFieldsMatrix(fresh_state)
    field_label = dm.iter_varying_fields(fresh_state)[0][0]
    material_cell = matrix.cell_widget(1, field_label)

    material_cell.value = "bad|name"

    assert spec.per_sample_values.get(1) is None
    assert material_cell.layout.border != ""


# ---------------------------------------------------------------------------
# peroTF catalog vs. Excel_creator: the labels in perotf_utils.process_specs must be the
# ones peroTF's sheet_experiment.py writes, the process list must be Excel_creator's, and
# the header row must exist when smart_databaser generates it with is_testing=False.
# ---------------------------------------------------------------------------

EXCEL_CREATOR_PROCESSES = [
    "Spin Coating",
    "Evaporation",
    "Co-Evaporation",
    "Seq-Evaporation",
    "Sputtering",
    "ALD",
    "Cleaning O2-Plasma",
    "Cleaning UV-Ozone",
    "Inkjet Printing",
    "Slot Die Coating",
    "Dip Coating",
    "Laser Scribing",
    "Close Space Sublimation",
    "Lamination",
    "Annealing",
    "Generic Process",
    "Multijunction Info",
]


def test_available_processes_mirror_excel_creator():
    assert dm.AVAILABLE_PROCESSES == ["Experiment Info", *EXCEL_CREATOR_PROCESSES]


def _full_config(process_type):
    from perotf_utils.process_specs import PROCESSES

    meta = PROCESSES[process_type]["meta"]
    config = dict(meta.get("config_defaults", {}))
    for key, _label, _min, _max in meta["numeric_config"]:
        config[key] = 2
    for key, _label in meta["boolean_config"]:
        config[key] = True
    return config


@pytest.mark.parametrize("process_type", EXCEL_CREATOR_PROCESSES)
def test_every_mapped_catalog_label_is_a_column_excel_creator_writes(fresh_state, process_type):
    """A path whose label the sheet never writes could never autofill anything. Checked
    over the default config and a full one (every count 2, every checkbox on), since e.g.
    Spin Coating names its rotation columns differently for one and several steps."""
    written: set[str] = set()
    for config in (dm.default_config_for(process_type), _full_config(process_type)):
        process = fresh_state.add_process(process_type, config=config)
        written |= dm.relevant_field_keys_for_process(process)
    mapped = set(dm.PROCESS_TYPE_FIELD_PATHS.get(process_type, {}))
    indexed_beyond_config = {key for key in mapped if any(ch.isdigit() for ch in key)}
    assert mapped - indexed_beyond_config <= written
    assert written, process_type


def test_experiment_info_mapped_labels_are_written(fresh_state):
    column_map = dm.rebuild_field_specs(fresh_state)
    written = {key for seq, key in column_map if seq == 0}
    assert set(dm.PROCESS_TYPE_FIELD_PATHS["Experiment Info"]) <= written


def test_excel_creator_header_row_written_when_not_testing(mods):
    """Regression test for sheet_experiment.add_experiment_sheet(..., is_testing=False):
    row 2 (the column labels) used to be written only in testing mode, which left
    smart_databaser with an empty column map. Row 3 and the Nomad ID formula stay
    testing-only."""
    from openpyxl import Workbook

    workbook = Workbook()
    mods["sheet_experiment"].add_experiment_sheet(
        workbook,
        [{"process": "Experiment Info"}, {"process": "Spin Coating", "config": {}}],
        is_testing=False,
    )
    ws = workbook.active
    row2 = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    assert row2[:6] == ["Date", "Project_Name", "Batch", "Subbatch", "Sample", "Nomad ID"]
    assert all(row2)
    assert "Material name" in row2
    assert all(ws.cell(row=3, column=c).value is None for c in range(1, ws.max_column + 1))


def test_excel_creator_testing_mode_still_writes_test_row_and_formula(mods):
    from openpyxl import Workbook

    workbook = Workbook()
    mods["sheet_experiment"].add_experiment_sheet(
        workbook, [{"process": "Experiment Info"}], is_testing={"any": "truthy config"}
    )
    ws = workbook.active
    assert ws["A2"].value == "Date"
    assert ws["A3"].value == "YYYYMMDD"
    assert ws["F3"].value.startswith('=CONCATENATE("KIT_"')


def test_no_atmospheric_checkbox_for_perotf(fresh_state):
    """Excel_creator has no atmospheric block, so the catalog leaves the key unset and no
    process row offers the checkbox."""
    assert dm.ATMOSPHERIC_CONFIG_KEY is None
    fresh_state.add_process("Spin Coating", config=dm.default_config_for("Spin Coating"))
    dm.rebuild_field_specs(fresh_state)
    builder = gc.ProcessSequenceBuilder(fresh_state, "url", "token", dm.NomadSessionCache())
    row = builder._process_rows[id(fresh_state.get_process(1))]
    descriptions = {w.description for w in row._checkbox_controls}
    assert "Add Atmospheric Values" not in descriptions
    assert {"Antisolvent", "Gas Quenching", "Vacuum Quenching"} <= descriptions


def test_batch_picker_lists_newest_batch_first_and_search_keeps_that_order(fresh_state):
    """The replicate/duplicate pickers sort like perotf_utils.batch_selection: newest
    YYYYMMDD date in the id first, ids without a date last."""
    cache = dm.NomadSessionCache()
    unsorted = [
        "KIT_JoDa_20250101_A_0",
        "KIT_JoDa_20260526_B_0",
        "no_date_batch",
        "KIT_HaGu_20251113_K16_0",
        "KIT_JoDa_20260526_A_0",
    ]
    with patch.object(dm, "get_batch_ids", return_value=unsorted):
        picker = gc.create_whole_experiment_template_picker(
            fresh_state, "url", "token", cache, on_change=lambda: None
        )
    search_field, selector, _load_button, _status = picker.children[2].children

    assert list(selector.options) == [
        "KIT_JoDa_20260526_A_0",
        "KIT_JoDa_20260526_B_0",
        "KIT_HaGu_20251113_K16_0",
        "KIT_JoDa_20250101_A_0",
        "no_date_batch",
    ]
    search_field.value = "joda"
    assert list(selector.options) == [
        "KIT_JoDa_20260526_A_0",
        "KIT_JoDa_20260526_B_0",
        "KIT_JoDa_20250101_A_0",
    ]


def test_initialize_ui_shows_the_beta_notice():
    with (
        patch.object(dm, "get_batch_ids", return_value=["B1"]),
        patch.object(dm, "get_all_uploads", return_value=[{"upload_id": "UP1"}]),
    ):
        main_interface = app_module.initialize_ui("url", "token")
    assert app_module.BETA_NOTICE_HTML in main_interface.children[0].value
    assert "Beta" in app_module.BETA_NOTICE_HTML
